"""End-to-end test: sample Asana + Adaptive raw extracts -> ServiceNow load files.

Fixtures mirror the real source structures (synthetic names/values):
  Asana 100-700   standard projects, pipeline portfolio, intake board, template
  Asana 900       HR portfolio tracker: each task is a project; 950 is 901's plan project
  Adaptive A1-A5  Nordic IT PMO projects with the tenant's C_ fields
"""
import shutil
from pathlib import Path

import pytest
import yaml

from spm_migration import cli
from spm_migration.common import read_csv

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture()
def run_dir(tmp_path):
    shutil.copytree(FIXTURES / "raw", tmp_path / "raw")
    settings = yaml.safe_load((ROOT / "config/settings.example.yaml").read_text())
    settings["paths"] = {k: str(tmp_path / k) for k in ("raw", "staging", "load", "reference")}
    settings["as_of_date"] = "2026-09-22"
    settings["asana"]["intake_project_gids"] = ["300"]
    settings["asana"]["tracker_project_gids"] = ["900"]
    settings["asana"]["plan_links_file"] = str(tmp_path / "plan_links.csv")
    (tmp_path / "plan_links.csv").write_text("tracker_item_gid,plan_project_gid\n901,950\n")
    ref = tmp_path / "reference"
    ref.mkdir()
    (ref / "sys_user.csv").write_text(
        "email,name\npat.pm@example.com,Pat Pm\ndev@example.com,Dev One\npm2@example.com,Pm Two\n"
        "pat.lee@example.org,Patricia Lee\n")
    (ref / "user_crosswalk.csv").write_text(
        "source_value,email\nPat Lee,pat.lee@example.org\nCasey Sponsor,casey.s@example.org\n")
    (tmp_path / "settings.yaml").write_text(yaml.safe_dump(settings))
    (tmp_path / "overrides.csv").write_text(
        "source_system,source_id,target_class,include_tasks,demand_link,note\n"
        "asana,600,demand,,,Owner says still an idea\n")
    args = ["--settings", str(tmp_path / "settings.yaml"), "--overrides", str(tmp_path / "overrides.csv"),
            "--rules", str(ROOT / "config/classification.yaml"),
            "--mapping", str(ROOT / "config/target_mapping.yaml"),
            "--value-maps", str(ROOT / "mapping/value_maps.csv")]
    assert cli.main(["run-all", *args]) == 0
    return tmp_path


def by(rows, key="correlation_id"):
    return {r[key]: r for r in rows}


def classified(run_dir):
    return {(r["source_system"], r["source_id"]): r
            for r in read_csv(run_dir / "staging/work_items_classified.csv")}


def test_classification_standard_asana(run_dir):
    rows = classified(run_dir)
    cls = {k: r["target_class"] for k, r in rows.items()}
    assert cls[("asana", "100")] == "project"
    assert cls[("asana", "200")] == "demand"             # pipeline portfolio
    assert cls[("asana", "300")] == "skip"               # intake board container
    assert cls[("asana", "301")] == "demand"             # intake task
    assert cls[("asana", "400")] == "skip"               # template
    assert cls[("asana", "500")] == "skip"               # old closed
    assert cls[("asana", "600")] == "demand"             # override beats empty-shell review rule
    assert rows[("asana", "600")]["classification_rule"] == "override"
    assert cls[("asana", "700")] == "project" and rows[("asana", "700")]["include_tasks"] == "false"


def test_classification_hr_tracker(run_dir):
    rows = classified(run_dir)
    rule = {k[1]: r["classification_rule"] for k, r in rows.items() if k[0] == "asana"}
    cls = {k[1]: r["target_class"] for k, r in rows.items() if k[0] == "asana"}
    assert cls["900"] == "skip" and cls["950"] == "skip"          # tracker + plan project are containers
    assert cls["901"] == "project" and cls["902"] == "project"
    assert "903" not in cls                                       # subtasks are WBS, not work items
    assert rule["905"] == "asana_not_started_without_pm_needs_review"
    assert rule["906"] == "completed_before_2025_is_skipped"
    assert rule["907"] == "cancelled_work_is_skipped"
    assert rule["911"] == "recently_completed_projects_header_only"
    assert rule["912"] == "asana_tracker_item_without_status_or_dates_needs_review"
    assert rows[("asana", "901")]["portfolio"] == "Benefits & Well-Being"   # section (COE) -> portfolio


def test_classification_adaptive(run_dir):
    rows = classified(run_dir)
    rule = {k[1]: r["classification_rule"] for k, r in rows.items() if k[0] == "adaptive"}
    assert rule["/Project/A1"] == "active_adaptive_projects"
    assert rule["/Project/A2"] == "adaptive_requested_or_draft_is_demand"
    assert rule["/Project/A3"] == "cancelled_work_is_skipped"
    assert rule["/Project/A4"] == "adaptive_active_without_delivery_handoff_needs_review"
    assert rule["/Project/A5"] == "recently_completed_projects_header_only"
    review = {r["source_id"] for r in read_csv(run_dir / "staging/classification_review.csv")}
    assert review == {"905", "912", "/Project/A4"}


def test_demand_load_file(run_dir):
    demands = by(read_csv(run_dir / "load/01_dmn_demand.csv"))
    crm = demands["ASANA:301"]
    assert crm["state"] == "submitted" and crm["demand_type"] == "hr"
    assert crm["opened_by"] == "migration.unassigned@example.com"   # requester not in sys_user
    assert crm["business_case"] == "Save $1M" and crm["priority"] == "3"
    assert crm["business_category"] == "Human Resources"
    a2 = demands["ADAPTIVE:/Project/A2"]
    assert a2["demand_type"] == "new_project" and a2["state"] == "submitted"
    assert a2["business_category"] == "Construction Renovation Relocation"   # demand choice list
    assert a2["sites"] == "Lorain,Lima"
    assert a2["opened_by"] == "casey.s@example.org"                   # falls back to business sponsor
    assert a2["business_owner"] == "casey.s@example.org"            # first resolvable of "A, B"
    assert "Business sponsor(s): Unknown Person, Casey Sponsor" in a2["description"]
    assert "estimate_type" not in a2                               # demand stays out-of-box + form custom fields
    assert "Estimate type: SAT - Detailed Estimate" in a2["description"]
    assert "Request received: 2026-07-01" in a2["description"]


def test_project_load_file(run_dir):
    projects = by(read_csv(run_dir / "load/02_pm_project.csv"))
    erp = projects["ASANA:100"]
    assert erp["state"] == "2" and erp["project_manager"] == "pat.pm@example.com"
    assert erp["description"] == "Upgrade the ERP"

    hr = projects["ASANA:901"]
    assert hr["project_manager"] == "pat.lee@example.org"            # "Pat Lee" via crosswalk
    assert hr["description"].startswith("· Vendor RFP – in progress")  # mojibake repaired
    assert "HR COEs engaged: Talent & Culture, HRTS" in hr["description"]
    assert hr["primary_portfolio"] == "Benefits & Well-Being"          # section / COE -> portfolio
    assert hr["primary_program"] == ""
    assert hr["business_unit"] == "BSMH" and hr["impacted_business_units"] == "BSMH,RSFH"
    assert hr["expense_type"] == "opex" and hr["investment_class"] == "change"
    assert projects["ASANA:902"]["investment_class"] == "run"        # Annual Program
    assert projects["ASANA:902"]["state"] == "-5"                    # Not Started -> Pending

    a1 = projects["ADAPTIVE:/Project/A1"]
    assert a1["legacy_id"] == "P-00001" and a1["cn"] == "48213"
    assert a1["business_owner"] == "casey.s@example.org"
    assert a1["business_category"] == "infrastructure"
    assert a1["sites"] == "Cincinnati,Toledo,Kentucky"
    assert a1["funding_source"] == "CIN (Cincinnati) Capital" and a1["expense_type"] == "capex"
    assert a1["ready_for_delivery"] == "2025-10-01" and a1["next_go_live_date"] == "2026-11-15"
    assert projects["ADAPTIVE:/Project/A5"]["state"] == "3"


def test_task_load_file(run_dir):
    tasks = read_csv(run_dir / "load/03_pm_project_task.csv")
    t = by(tasks)
    assert t["ASANA:101"]["parent_correlation_id"] == "ASANA:section-s1"   # sections as phases
    assert t["ASANA:102"]["parent_correlation_id"] == "ASANA:101" and t["ASANA:102"]["level"] == "3"
    assert t["ASANA:102"]["start_date"] == "2026-01-29"                     # due-only task -> start = due
    assert t["ASANA:104"]["milestone"] == "true"
    # HR tracker subtasks -> WBS of the tracker item
    assert t["ASANA:903"]["project_correlation_id"] == "ASANA:902"
    assert t["ASANA:903"]["parent_correlation_id"] == "ASANA:902"
    # plan project 950 -> WBS of tracker item 901
    m1 = t["ASANA:951"]
    assert m1["project_correlation_id"] == "ASANA:901" and m1["parent_correlation_id"] == "ASANA:901"
    assert m1["milestone"] == "true" and m1["percent_complete"] == "50.0" and m1["state"] == "2"
    assert m1["additional_assignee_list"] == "dev@example.com"   # assignee + unknown followers dropped
    assert "ASANA:section-u1" not in t                           # 'Untitled section' is not a phase
    assert t["ASANA:954"]["parent_correlation_id"] == "ASANA:section-p2"
    assert t["ADAPTIVE:/Task/T2"]["parent_correlation_id"] == "ADAPTIVE:/Task/T1"
    assert t["ADAPTIVE:/Milestone/M1"]["milestone"] == "true"
    assert "ASANA:701" not in t                                   # header-only project
    levels = [int(r["level"]) for r in tasks]
    assert levels == sorted(levels)                               # parents before children
    excluded = {r["source_id"]: r["exclusion_reason"] for r in read_csv(run_dir / "load/excluded_tasks.csv")}
    assert excluded["701"] == "project loaded header-only"


def test_dependencies_status_and_reports(run_dir):
    deps = read_csv(run_dir / "load/04_planned_task_rel_planned_task.csv")
    pairs = {(d["predecessor_correlation_id"], d["successor_correlation_id"]): d for d in deps}
    assert ("ASANA:101", "ASANA:103") in pairs and ("ASANA:951", "ASANA:954") in pairs
    assert pairs[("ADAPTIVE:/Task/T1", "ADAPTIVE:/Milestone/M1")]["lag"] == "2.0"
    assert len(read_csv(run_dir / "load/excluded_dependencies.csv")) == 1   # 999 not extracted

    status = {(s["project_correlation_id"], s["as_on"]): s for s in read_csv(run_dir / "load/05_project_status.csv")}
    assert status[("ASANA:100", "2026-09-01")]["overall_health"] == "yellow"
    a1 = status[("ADAPTIVE:/Project/A1", "2026-09-10")]                     # six Adaptive health fields
    assert (a1["overall_health"], a1["schedule"], a1["cost"], a1["resources"], a1["scope"], a1["risk_health"]) == \
        ("yellow", "yellow", "green", "green", "green", "red")
    assert a1["comments"] == "Wave 1 complete; vendor delay on wave 2"

    users = {u["source_value"]: u for u in read_csv(run_dir / "load/users_referenced.csv")}
    assert users["requester@example.com"]["matched"] == "NO"
    assert users["pat lee"]["resolved_email"] == "pat.lee@example.org"
    assert "tbd" not in users                                    # TBD / N/A treated as blank
    assert [m["source_id"] for m in read_csv(run_dir / "staging/multi_homed_tasks.csv")] == ["103"]
    unmapped = {(u["map_name"], u["source_value"]) for u in read_csv(run_dir / "load/unmapped_values.csv")}
    assert unmapped == {("business_category_project", "BSMH")}      # only the open D17 gap


def test_validate_fields(run_dir):
    ref = run_dir / "reference"
    (ref / "sys_dictionary.csv").write_text(
        "name,element\n"
        "task,correlation_id\ntask,correlation_display\ntask,short_description\ntask,description\n"
        "task,assigned_to\ntask,opened_by\ntask,state\ntask,priority\ntask,due_date\n"
        "planned_task,start_date\nplanned_task,end_date\nplanned_task,work_start\nplanned_task,work_end\n"
        "planned_task,percent_complete\nplanned_task,time_constraint\nplanned_task,wbs_order\n"
        "planned_task,milestone\ndmn_demand,type\ndmn_demand,category\ndmn_demand,portfolio\n"
        "pm_project,project_manager\npm_project,phase\npm_project,primary_portfolio\n"
        "project_status,as_on\nproject_status,overall_health\n")
    (ref / "sys_choice.csv").write_text("name,element,value\nproject_status,overall_health,green\n"
                                        "project_status,overall_health,red\n")
    args = ["--settings", str(run_dir / "settings.yaml"), "--mapping", str(ROOT / "config/target_mapping.yaml"),
            "--value-maps", str(ROOT / "mapping/value_maps.csv")]
    assert cli.main(["validate-fields", *args]) == 1
    issues = read_csv(run_dir / "staging/field_validation.csv")
    errors = {(i["table"], i["column"]) for i in issues if i["severity"] == "ERROR"}
    assert ("project_status", "comments") in errors                  # not in dictionary
    assert ("project_status", "overall_health") in errors            # yellow not in choices
    assert ("pm_project", "short_description") not in errors         # inherited from task


def test_reconcile(run_dir):
    (run_dir / "reference/sn_pm_project.csv").write_text(
        "number,correlation_id\nPRJ001,ASANA:100\nPRJ002,ASANA:777\n")
    args = ["--settings", str(run_dir / "settings.yaml"), "--mapping", str(ROOT / "config/target_mapping.yaml")]
    assert cli.main(["reconcile", *args]) == 0
    summary = {r["table"]: r for r in read_csv(run_dir / "load/reconciliation_summary.csv")}
    assert summary["pm_project"]["unexpected"] == "1"
    detail = read_csv(run_dir / "load/reconciliation_detail.csv")
    assert {"table": "pm_project", "correlation_id": "ASANA:777",
            "problem": "in ServiceNow but not in load file"} in detail


def test_helpers():
    from spm_migration.common import parse_percent, repair_text, split_multi
    assert repair_text("¬∑¬†¬†Feb 5-19 ‚Äì Pulse") == "· Feb 5-19 – Pulse"
    assert split_multi("BSMH, RSFH") == ["BSMH", "RSFH"] and split_multi(["a", "b"]) == ["a", "b"]
    assert parse_percent(0.5) == 50 and parse_percent("50%") == 50 and parse_percent(75) == 75


def test_mapping_csvs_are_well_formed():
    import csv
    for path in (ROOT / "mapping").glob("*.csv"):
        rows = list(csv.reader(path.open(encoding="utf-8")))
        assert all(len(r) == len(rows[0]) for r in rows), f"{path.name} has rows with the wrong column count"


def test_task_rollups_to_parent_project(run_dir):
    """Asana task mapping: COE -> parent.portfolio, Entity/Organization -> parent.u_business_category."""
    projects = by(read_csv(run_dir / "load/02_pm_project.csv"))
    # 700 has no portfolio of its own: its task's COE 'Learning' fills it (value-mapped)
    assert projects["ASANA:700"]["primary_portfolio"] == "Culture & Learning"
    # 901's section is its portfolio; task COEs disagree -> section wins, conflict reported
    assert projects["ASANA:901"]["primary_portfolio"] == "Benefits & Well-Being"
    assert projects["ASANA:901"]["business_category"] == "rsfh"        # 'BSMH, RSFH' -> first mappable
    conflicts = {(c["source_id"], c["field"]): c for c in read_csv(run_dir / "staging/rollup_conflicts.csv")}
    assert conflicts[("901", "portfolio")]["used"] == "Benefits & Well-Being"
    assert ("700", "portfolio") not in conflicts
    unmapped = {(u["map_name"], u["source_value"]) for u in read_csv(run_dir / "load/unmapped_values.csv")}
    assert ("business_category_project", "BSMH") in unmapped                  # D17: no BSMH category


def test_transform_map_spec(tmp_path):
    from spm_migration.servicenow import transform_map_spec
    rows = transform_map_spec(ROOT / "config/target_mapping.yaml", tmp_path / "spec.csv")
    assert all("u_u_" not in r["source_column"] for r in rows)                # no double prefix
    spec = {(r["source_system"], r["target_table"], r["source_column"]): r for r in rows}
    assert spec[("asana", "pm_project", "u_correlation_id")]["coalesce"] == "Yes"
    assert spec[("adaptive", "pm_project", "u_cn")]["import_set_table"] == "u_imp_adaptive_project"
    assert ("asana", "pm_project", "u_cn") not in spec                        # Adaptive-only field
    owner = spec[("asana", "pm_project", "u_business_owner")]
    assert (owner["field_kind"], owner["referenced_value_field"], owner["choice_action"],
            owner["target_field_status"]) == ("Reference sys_user", "email", "ignore", "exists (custom)")
    assert spec[("adaptive", "pm_project", "u_sites")]["field_kind"] == "List"
    assert spec[("adaptive", "pm_project", "u_next_go_live_date")]["target_field_status"] == "PROPOSED - create before load"
    assert spec[("asana", "pm_project_task", "u_parent_correlation_id")]["field_kind"] == "Helper"
    assert spec[("adaptive", "planned_task_rel_planned_task", "u_predecessor_correlation_id")]["target_field"] == "parent"
    assert spec[("adaptive", "pm_project", "u_state")]["choice_action"] == "reject"


def test_push_dry_run_and_execute(run_dir, monkeypatch):
    from spm_migration import servicenow
    settings = yaml.safe_load((run_dir / "settings.yaml").read_text())
    summary = servicenow.push(settings, ROOT / "config/target_mapping.yaml", only=["pm_project"])
    assert [s["import_set_table"] for s in summary] == ["u_imp_asana_project", "u_imp_adaptive_project"]
    assert all(s["dry_run"] == s["rows"] > 0 for s in summary)

    sent = []

    class FakeResp:
        status_code = 201
        headers = {}

        def json(self):
            return {"result": [{"status": "inserted", "sys_id": "abc123"}]}

    def fake_post(self, url, json=None, timeout=None):
        sent.append((url, json))
        return FakeResp()

    monkeypatch.setattr("requests.Session.post", fake_post)
    settings["servicenow"].update(instance_url="https://x.service-now.com", username="u", password="p")
    summary = servicenow.push(settings, ROOT / "config/target_mapping.yaml", only=["pm_project"],
                              execute=True, limit=2, sources=("adaptive",))
    assert summary[0]["inserted"] == 2
    assert sent[0][0] == "https://x.service-now.com/api/now/import/u_imp_adaptive_project"
    assert "u_correlation_id" in sent[0][1] and "u_short_description" in sent[0][1]
    assert "u_cn" in sent[0][1] and not any(k.startswith("u_u_") for k in sent[0][1])
    results = read_csv(run_dir / "load/push_results.csv")
    assert results[0]["target_sys_id"] == "abc123"


def test_load_files_split_by_source_without_u_prefix(run_dir):
    header = (run_dir / "load/adaptive/02_pm_project.csv").read_text().splitlines()[0].split(",")
    assert "cn" in header and "business_owner" in header and "short_description" in header
    assert not any(h.startswith("u_") for h in header)
    asana = read_csv(run_dir / "load/asana/02_pm_project.csv")
    adaptive = read_csv(run_dir / "load/adaptive/02_pm_project.csv")
    assert asana and all(r["correlation_id"].startswith("ASANA:") for r in asana)
    assert adaptive and all(r["correlation_id"].startswith("ADAPTIVE:") for r in adaptive)
    assert len(asana) + len(adaptive) == len(read_csv(run_dir / "load/02_pm_project.csv"))


def test_user_stories_generation(tmp_path):
    import openpyxl
    from spm_migration.servicenow import transform_map_spec
    from spm_migration.stories import build
    transform_map_spec(ROOT / "config/target_mapping.yaml", tmp_path / "spec.csv")
    xlsx, md = build(ROOT / "config/user_stories.yaml", tmp_path / "spec.csv", ROOT / "mapping/decisions.csv",
                     tmp_path / "stories.xlsx", tmp_path / "stories.md")
    wb = openpyxl.load_workbook(xlsx)
    assert wb.sheetnames == ["Read Me", "Epics", "Stories", "Field Maps", "Open Items"]
    stories = list(wb["Stories"].iter_rows(min_row=2, values_only=True))
    keys = [s[0] for s in stories]
    assert keys[:4] == ["MIG-01", "MIG-02", "MIG-03", "MIG-04"] and "MIG-09" in keys
    mig09 = next(s for s in stories if s[0] == "MIG-09")
    assert "u_imp_adaptive_demand" in mig09[3] and "u_business_owner" in mig09[3]
    text = md.read_text()
    assert "`u_u_" not in text                                        # no double-prefixed columns
    fm = list(wb["Field Maps"].iter_rows(min_row=2, values_only=True))
    assert fm and not any(str(r[4]).startswith("u_u_") for r in fm)
    assert "`u_cn` | `u_cn`" in text


def test_export_files_route(tmp_path):
    """No API access: Asana CSV exports + Adaptive Excel exports -> same load files."""
    from openpyxl import Workbook
    exp = tmp_path / "exports"
    (exp / "asana").mkdir(parents=True)
    (exp / "adaptive").mkdir()
    (exp / "asana/tracker.csv").write_text(
        "Task ID,Name,Section/Column,Assignee,Start Date,Due Date,Completed At,Notes,Parent task,"
        "Blocked By (Dependencies),Project Status,PM Assigned,Project Type,Entity/Organization\n"
        "111,Benefits Refresh,Benefits & Well-Being,,2026-01-05,2026-12-31,,Vendor RFP,,,In Progress,Pat Lee,Project,RSFH\n"
        "112,Vendor A,,,2026-02-01,2026-03-01,,,Benefits Refresh,,,,,\n"
        "113,Old Thing,Labor,,2024-01-01,2024-02-01,,,,,Cancelled,Pat Lee,Project,BSMH\n")
    (exp / "asana/plan.csv").write_text(
        "Task ID,Task Name,Section,Assigned To,Due Date,Task Status,% Complete,Parent task,Blocked By (Dependencies),COE\n"
        "201,Initiation Milestone 1,Project Initiation,Pat Lee,2026-08-26,In Progress,0.5,,,Benefits\n"
        "202,Kickoff Meeting,,,,,,Initiation Milestone 1,,\n"
        "203,Build Task 1,Build,,2026-10-01,Not Started,0,,Initiation Milestone 1,Benefits\n")
    wb = Workbook()
    ws = wb.active
    ws.append(["SYSID", "Name", "Manager", "State", "Start Date", "Planned Finish", "CN#", "Region",
               "Business Sponsor", "Health", "Ready For Delivery"])
    ws.append(["P-1", "DC Exit", "pm2@example.com", "Active", "2026-01-01", "2027-03-31", "00123",
               "Cincinnati; Toledo", "Casey Sponsor", "In Trouble", "2025-10-01"])
    ws.append(["P-2", "Clinic Build", "Sam Doe", "Requested", "", "", "", "Lima", "", "", ""])
    wb.save(exp / "adaptive/Adaptive_Projects.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["SYSID", "Name", "Project ID", "Parent ID", "State", "Start Date", "Due Date"])
    ws.append(["T-1", "Wave 1", "P-1", "P-1", "Active", "2026-02-01", "2026-04-30"])
    ws.append(["T-2", "Wave 1a", "P-1", "T-1", "Completed", "2026-02-01", "2026-03-01"])
    wb.save(exp / "adaptive/Adaptive_Tasks.xlsx")

    settings = yaml.safe_load((ROOT / "config/settings.example.yaml").read_text())
    settings["paths"] = {k: str(tmp_path / k) for k in ("raw", "staging", "load", "reference", "exports")}
    settings["as_of_date"] = "2026-09-24"
    settings["asana"]["exports"] = [{"file": "tracker.csv", "role": "tracker"},
                                    {"file": "plan.csv", "role": "plan", "tracker_item": "Benefits Refresh"}]
    settings["asana"]["plan_links_file"] = str(tmp_path / "none.csv")
    for entity in ("Milestone", "Link"):
        settings["adaptive"]["exports"].pop(entity)
        settings["adaptive"]["entities"].pop(entity)
    (tmp_path / "reference").mkdir()
    (tmp_path / "reference/user_crosswalk.csv").write_text(
        "source_value,email\nPat Lee,pat.lee@example.org\nCasey Sponsor,casey.s@example.org\n")
    (tmp_path / "settings.yaml").write_text(yaml.safe_dump(settings))
    base = ["--settings", str(tmp_path / "settings.yaml"), "--rules", str(ROOT / "config/classification.yaml"),
            "--mapping", str(ROOT / "config/target_mapping.yaml"), "--value-maps", str(ROOT / "mapping/value_maps.csv"),
            "--overrides", str(tmp_path / "none.csv")]
    assert cli.main(["import-asana-exports", *base]) == 0
    assert cli.main(["import-adaptive-exports", *base]) == 0
    assert read_csv(tmp_path / "raw/asana/export_warnings.csv") == []
    assert cli.main(["run-all", *base]) == 0

    projects = by(read_csv(tmp_path / "load/asana/02_pm_project.csv"))
    assert set(projects) == {"ASANA:111"}                                   # 113 cancelled -> skipped
    assert projects["ASANA:111"]["project_manager"] == "pat.lee@example.org"
    assert projects["ASANA:111"]["primary_portfolio"] == "Benefits & Well-Being"
    tasks = by(read_csv(tmp_path / "load/asana/03_pm_project_task.csv"))
    assert tasks["ASANA:112"]["parent_correlation_id"] == "ASANA:111"       # tracker subtask
    assert tasks["ASANA:202"]["parent_correlation_id"] == "ASANA:201"       # plan subtask by parent name
    assert tasks["ASANA:201"]["project_correlation_id"] == "ASANA:111"      # plan -> tracker item
    assert tasks["ASANA:201"]["milestone"] == "true" and tasks["ASANA:201"]["percent_complete"] == "50.0"
    deps = read_csv(tmp_path / "load/asana/04_planned_task_rel_planned_task.csv")
    assert [(d["predecessor_correlation_id"], d["successor_correlation_id"]) for d in deps] == [("ASANA:201", "ASANA:203")]

    a_proj = by(read_csv(tmp_path / "load/adaptive/02_pm_project.csv"))
    a1 = a_proj["ADAPTIVE:/Project/P-1"]
    assert a1["cn"] == "00123" and a1["sites"] == "Cincinnati,Toledo" and a1["business_owner"] == "casey.s@example.org"
    assert a1["ready_for_delivery"] == "2025-10-01"
    demands = by(read_csv(tmp_path / "load/adaptive/01_dmn_demand.csv"))
    assert set(demands) == {"ADAPTIVE:/Project/P-2"}                        # Requested -> demand
    a_tasks = by(read_csv(tmp_path / "load/adaptive/03_pm_project_task.csv"))
    assert a_tasks["ADAPTIVE:/Task/T-2"]["parent_correlation_id"] == "ADAPTIVE:/Task/T-1"
    assert a_tasks["ADAPTIVE:/Task/T-1"]["parent_correlation_id"] == "ADAPTIVE:/Project/P-1"
    status = read_csv(tmp_path / "load/adaptive/05_project_status.csv")
    assert status[0]["overall_health"] == "red"
