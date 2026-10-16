"""Build docs/rm_story_drafts_ready.xlsx: the Draft migration stories, kept short.
Each story = a few plain steps + the script file to use. Detail lives in the scripts (servicenow/),
not in the story.
"""
import html
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "docs" / "rm_story_drafts_ready.xlsx"

# number, short description, points, assign now, criteria
DRAFTS = [
    ("STRY0067504", "Set up the Adaptive and Asana API connection settings", 1, "Yes", [
        "These system properties exist in Dev, Test and Prod: x_mig.adaptive.url, x_mig.adaptive.key (password2), "
        "x_mig.asana.token (password2), x_mig.asana.project_gid.",
        "The keys are the svc_sn_migration service account keys.",
        "servicenow/test_connection.js, run in Scripts - Background, shows Status: 200 for Adaptive and Asana in each instance.",
        "No Outbound REST Message is needed.",
    ]),
    ("STRY0067506", "Create the Adaptive Data Sources for project-related records", 8, "Yes", [
        "servicenow/discovery_background_script.js has been run once, and its output is attached to this story.",
        "One Data Source (Type: Custom (Load by Script)) exists for each: Tasks, Milestones, Risks, Issues, "
        "Lessons Learned, Resources, Time.",
        "Each uses servicenow/adaptive_related_data_source.js. Only the settings block at the top changes, using the "
        "names from the Discovery output.",
        "Each import table has the columns Correlation ID, Parent Correlation ID and Project Correlation ID.",
        "Load All Records runs with no errors for the pilot projects.",
    ]),
    ("STRY0066701", "Create status reports from Adaptive project health", 3, "Yes", [
        "A new transform map turns the Adaptive Projects import into status reports (project_status), set up as "
        "described at the top of servicenow/status_report_transform.js.",
        "Health colours: On Target / On Plan = Green, Needs Attention = Yellow, In Trouble = Red.",
        "Update Notes go into the status report comments.",
        "Each pilot project has one status report, and running it again adds no new ones.",
    ]),
    ("STRY0066702", "Add people and hours worked to project tasks (Adaptive)", 8, "Yes", [
        "Adaptive resource assignments are loaded: the first person becomes Assigned to; others go into Additional "
        "assignee list.",
        "Adaptive time is loaded as total hours per task into the task's actual effort.",
        "These loads only update existing tasks; they never create new ones. Scripts: "
        "servicenow/resources_time_transform.js.",
        "Rate models, weekly time cards and resource plans are out of scope.",
        "5 pilot tasks show the same people and hours as Adaptive.",
    ]),
    ("STRY0066723", "Copy attachments from Asana and Adaptive", 5, "Yes - run last", [
        "servicenow/attachments_fix_script.js is run as a Fix Script after all other loads.",
        "It is run first with DRY_RUN = true, and the log reviewed. Then it is run with DRY_RUN = false.",
        "Files stored in Google Drive, Box, Dropbox or OneDrive are added to the record as a link in a work note.",
        "Running it again does not create duplicate files.",
        "10 spot-checked files open correctly in ServiceNow.",
    ]),
    ("STRY0066721", "Create cost plans from Adaptive budgets", 6, "HOLD", [
        "On hold until the BA confirms (1) where budgets live in Adaptive and (2) cost plans or demand budget totals "
        "(not both).",
    ]),
]

# number, state, one-line fix
IN_FLIGHT = [
    ("STRY0067932", "Ready", "Use servicenow/asana_data_source.js in a Data Source instead of a Scheduled Job. No "
     "Outbound REST Message needed. Project GID = 1208928603143252. Table = u_import_asana_hr_project_portfolio."),
    ("STRY0066697", "WIP", "Scripts must leave fields blank when no match (not 'undefined'). Add Project Description -> "
     "Detailed Description. Business Sponsor -> Business Owner."),
    ("STRY0066699", "WIP", "Feed from the Tasks and Milestones Data Sources (STRY0067506). Run the transform twice so "
     "every task finds its parent."),
    ("STRY0066695", "WIP", "On hold until the Demand vs Project rule is known. Total Budget must not go to financial_return."),
    ("STRY0066715", "WIP", "Match records on correlation_id (not parent + title)."),
    ("STRY0066718", "WIP", "Match records on correlation_id (not parent + identifier)."),
    ("STRY0067099", "WIP", "correlation_id and the project lookup must use the \"ADAPTIVE:\" prefix, or lessons won't "
     "find their project."),
    ("STRY0065263", "WIP", "Load from the Asana API, not CSV. Mappings wait on the 12 Asana questions."),
    ("STRY0065275", "WIP", "Load from the Asana API (asana_data_source.js, MODE = 'tasks'). Parent comes from Parent "
     "Correlation ID, not the task name."),
]

ORDER = ["STRY0067504 Connection settings", "STRY0066697 Projects (fix + rerun)", "STRY0067932 / STRY0065263 Asana projects",
         "STRY0067506 Adaptive Data Sources", "STRY0066699 / STRY0065275 Tasks", "STRY0066701 Status reports",
         "STRY0066718 / STRY0066715 / STRY0067099 Risks, issues, lessons", "STRY0066702 People and hours",
         "STRY0066695 Demands (when unblocked)", "STRY0066721 Cost plans (when unblocked)", "STRY0066723 Attachments (last)"]


def criteria_text(items):
    return "This story is complete when:\n" + "\n".join("- " + i for i in items)


def criteria_html(items):
    lis = "".join(f"<li>{html.escape(i)}</li>" for i in items)
    return f'<p>This story is complete when:</p>\n<ul style="list-style-position: inside;">{lis}</ul>'


def sheet(wb, title, head, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet(wb, "Drafts", ["Number", "Short description", "Points", "Assign now?", "Acceptance criteria",
                              "Acceptance criteria (HTML - paste into ServiceNow)"],
               [(n, t, p, a, criteria_text(c), criteria_html(c)) for n, t, p, a, c in DRAFTS],
               [13, 44, 7, 14, 90, 30])
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 4).fill = PatternFill("solid", fgColor="FCE4D6" if ws.cell(r, 4).value == "HOLD" else "E2EFDA")
    sheet(wb, "In-flight fixes", ["Number", "State", "Fix"], IN_FLIGHT, [13, 10, 110])
    sheet(wb, "Build order", ["Step", "Story"], [(i + 1, s) for i, s in enumerate(ORDER)], [6, 70])
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
