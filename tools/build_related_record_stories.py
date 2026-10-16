"""Build docs/rm_story_related_records.xlsx: API-based stories for the records that hang off a project
(work plans, status reports, resource assignments, time, lessons learned, budgets, attachments).

Sheet "Stories": plain-text acceptance criteria + an HTML copy to paste into ServiceNow.
Sheet "Run order": the order loads must run in. Sheet "Decisions": open items the stories point to.
"""
import html
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "rm_story_related_records.xlsx"
P, UL, TABLE, CODE = "p", "ul", "table", "code"

SHARED = ("Built on what is proven in Dev: system properties x_mig.adaptive.url / x_mig.adaptive.key / "
          "x_mig.asana.token / x_mig.asana.project_gid, and the 'Adaptive Projects - API' Data Source "
          "(PRJ0010675 loaded from the API). Keys: correlation_id = \"ADAPTIVE:\" + Adaptive id "
          "(e.g. ADAPTIVE:/Project/5afh...) or \"ASANA:\" + Asana gid. Scripts are in the repo folder servicenow/.")

STAGING_COLS = ("Each staging table gets three String(100) columns if missing: Correlation ID, "
                "Parent Correlation ID, Project Correlation ID (same way Correlation ID was added to "
                "u_import_adaptive_projects).")

STORIES = [
    dict(num="NEW-10", pts=2, wave="A - now", title="API Discovery: confirm Adaptive entities/fields and Asana content",
         body=[
             (P, SHARED),
             (P, "Purpose: the related-record Data Sources need exact Adaptive entity and field names. One script "
                 "prints them, so nothing is guessed."),
             (UL, [
                 "Run servicenow/discovery_background_script.js in Scripts - Background (Dev), after at least one "
                 "Adaptive project is loaded.",
                 "Download the output and attach it to this story.",
                 "From the output, record on this story: the entity names for tasks, milestones, issues, risks, lessons "
                 "learned, resource links, timesheets and attachments; and the field on each that points at the "
                 "project (usually Project; lessons: C_RelatedWorkItem).",
                 "From the Asana section, record: whether portfolio items have subtasks, attachments, and "
                 "actual_time_minutes (time tracking).",
                 "Done when every [CONFIRM] value in STRY0067506, NEW-13 and STRY0066723 is filled in.",
             ]),
         ]),
    dict(num="STRY0067504", pts=1, wave="A - now", title="Create Outbound REST Message for AdaptiveWork (Clarizen) API",
         body=[
             (P, "Status: connection proven in Dev (POST /data/query and GET /metadata/describeEntities return 200)."),
             (P, "This story is complete when:"),
             (UL, [
                 "Base URL is the system property x_mig.adaptive.url = https://api2.clarizen.com/v2.0/services; "
                 "key in x_mig.adaptive.key (type password2). Asana: x_mig.asana.token (password2), "
                 "x_mig.asana.project_gid = HR Project Portfolio 2026.",
                 "The keys are switched to the svc_sn_migration service account keys (Adaptive and Asana) once issued; "
                 "only the property values change.",
                 "The same four properties exist in Test and Prod (values are entered per instance; they do not move "
                 "with code).",
                 "Field metadata is read with GET (not POST) /metadata/describeEntities?typeNames=<Entity>.",
             ]),
         ]),
    dict(num="STRY0067506", pts=8, wave="B", title="Create API Data Sources for Adaptive related records",
         body=[
             (P, SHARED),
             (P, STAGING_COLS),
             (P, "Pattern: copy servicenow/adaptive_related_data_source.js into a new Data Source "
                 "(Type: Custom (Load by Script)) and change only its CONFIG block. It loads records ONLY for projects "
                 "already in pm_project with an ADAPTIVE: correlation_id, 20 projects per query, following paging."),
             (TABLE, ["Data Source", "Staging table", "ENTITY", "PROJECT_FIELD", "Notes"], [
                 ["Adaptive Tasks - API", "u_import_adaptive_project_tasks", "Task", "Project",
                  "Parent Correlation ID = @Parent, Project Correlation ID = @Project, Milestone = '=false'"],
                 ["Adaptive Milestones - API", "u_import_adaptive_project_tasks", "Milestone", "Project",
                  "Same MAP; Milestone = '=true'"],
                 ["Adaptive Risks - API", "u_import_adaptive_risks", "Risk [CONFIRM]", "[CONFIRM]",
                  "MAP labels = the staging table's column labels"],
                 ["Adaptive Issues - API", "u_import_adaptive_issues", "Issue [CONFIRM]", "[CONFIRM]",
                  "MAP labels = the staging table's column labels"],
                 ["Adaptive Lessons - API", "u_import_adaptive_lessons (new)", "C_PMOLessonsLearned",
                  "C_RelatedWorkItem", "Preset in the script comments"],
                 ["Adaptive Resources - API", "u_import_adaptive_resources (new)", "[CONFIRM] e.g. RegularResourceLink",
                  "[CONFIRM]", "Columns: Correlation ID, Project Correlation ID (the work item), Resource Email, Units/Hours"],
                 ["Adaptive Time - API", "u_import_adaptive_time (new)", "[CONFIRM] e.g. Timesheet", "Project",
                  "Columns: Correlation ID, Project Correlation ID (@WorkItem), Reported By Email, Reported Date, Duration"],
             ]),
             (P, "This story is complete when:"),
             (UL, [
                 "Each Data Source above exists, and 'Load All Records' completes with no errors in the import log.",
                 "The import log line 'rows loaded: N (for M projects)' matches a count in Adaptive for the pilot projects.",
                 "Every staging row has Correlation ID and Project Correlation ID filled; task rows also Parent "
                 "Correlation ID.",
                 "Reference values arrive as plain text (e.g. 'Active', not {\"id\":\"/State/Active\"}); dates as yyyy-MM-dd.",
                 "Data Sources are not scheduled; they are run by hand in Run order (see Run order tab).",
             ]),
         ]),
    dict(num="NEW-11", pts=5, wave="B", title="Create API Data Sources for Asana HR portfolio (projects and work plans)",
         body=[
             (P, SHARED),
             (P, "Replaces the CSV loads for u_import_asana_hr_project_portfolio and u_import_asana_project_tasks. "
                 "The existing transform maps stay; the script fills the same columns (matched by label), and fills "
                 "any other column whose label equals an Asana custom field name (Project Status, PM Assigned, "
                 "Funding, Entity/Organization, ...)."),
             (UL, [
                 "Add String(100) columns: Correlation ID (both tables); Parent Correlation ID and Project Correlation ID "
                 "(tasks table); optional Assignee Email and Completed.",
                 "Data Source 'Asana Portfolio - API' -> u_import_asana_hr_project_portfolio, script "
                 "servicenow/asana_data_source.js with MODE = 'projects'.",
                 "Data Source 'Asana Work Plan - API' -> u_import_asana_project_tasks, same script with MODE = 'tasks' "
                 "(walks subtasks at every level).",
                 "PILOT list holds the pilot item names; [] loads every item. Items with Project Status Completed or "
                 "Cancelled are skipped (agreed in-flight rule).",
                 "Both transform maps get a field map u_correlation_id -> correlation_id with Coalesce = true.",
                 "People: transform maps match sys_user on email (Assignee Email column) and fall back to name.",
                 "Done when the pilot items and their subtasks load, transform, and a second run gives 0 inserts.",
             ]),
         ]),
    dict(num="NEW-12", pts=5, wave="C", title="Work plans: load Adaptive and Asana tasks into pm_project_task",
         body=[
             (P, SHARED),
             (P, "Uses the existing maps 'Import Adaptive Project Tasks Transform Map' and 'Import Asana Project "
                 "Tasks Transform Map'. Feeds: STRY0067506 (Tasks + Milestones) and NEW-11 (Work Plan)."),
             (UL, [
                 "Field map u_correlation_id -> correlation_id, Coalesce = true (both maps).",
                 "onBefore script (below) sets parent and top_task from the staging columns; a row whose project is "
                 "missing is skipped with a log message.",
                 "Run business rules = on, so dates and percent complete roll up to the project.",
                 "Children can arrive before parents: run the transform twice. The second run attaches the rest "
                 "(coalesce makes it an update).",
                 "Field maps per the value tables below. Task state values for Adaptive are PROPOSED (decision D1).",
                 "After the load, open each pilot project's Project Tasks list: the hierarchy matches Adaptive/Asana and "
                 "the project's dates are re-checked against the source (ServiceNow recalculates project dates from tasks).",
             ]),
             (CODE, "(function runTransformScript(source, map, log, target) {\n"
                    "    function find(table, corr) {\n"
                    "        var g = new GlideRecord(table);\n"
                    "        return (corr && g.get('correlation_id', corr)) ? g.getUniqueValue() : '';\n"
                    "    }\n"
                    "    var project = find('pm_project', source.u_project_correlation_id + '');\n"
                    "    if (!project) { log.warn('No project for ' + source.u_correlation_id); ignore = true; return; }\n"
                    "    var parentCorr = source.u_parent_correlation_id + '';\n"
                    "    var parent = (parentCorr == source.u_project_correlation_id + '') ? project : find('planned_task', parentCorr);\n"
                    "    target.top_task = project;\n"
                    "    target.parent = parent || project;   // parent not loaded yet: second run fixes it\n"
                    "})(source, map, log, target);"),
             (TABLE, ["Source", "Source value", "pm_project_task"], [
                 ["Adaptive Name / Asana name", "", "short_description"],
                 ["StartDate / start_on", "", "start_date"],
                 ["DueDate / due_on", "", "end_date"],
                 ["PercentCompleted", "", "percent_complete"],
                 ["Milestone column", "true", "milestone = true"],
                 ["Adaptive State (PROPOSED, D1)", "Draft, Requested, On Hold", "Pending"],
                 ["", "Active", "Work in Progress"],
                 ["", "Completed", "Closed Complete"],
                 ["", "Cancelled", "Closed Incomplete"],
                 ["Asana Task Status (agreed)", "Not Started, On Hold, blank", "Pending"],
                 ["", "Started/Scoping", "Investigating"],
                 ["", "In Progress, Delayed", "Work in Progress"],
                 ["", "Completed (or completed = true)", "Closed Complete"],
                 ["", "Cancelled", "Closed Incomplete"],
                 ["Asana Notes", "", "work_notes"],
                 ["Asana Collaborators", "names", "additional_assignee_list (sys_user by name)"],
             ]),
         ]),
    dict(num="STRY0066701", pts=3, wave="C", title="Status reports: Adaptive project health -> project_status",
         body=[
             (P, SHARED),
             (P, "Adaptive keeps health on the project itself (Health, Schedule/Budget/Resource/Scope/Risk Health, "
                 "Update Notes), so no extra API call is needed: a SECOND transform map on the same staging table "
                 "u_import_adaptive_projects creates one status report per project. Asana: no status reports for HR "
                 "portfolio items (they are tasks); not in scope."),
             (UL, [
                 "System property x_mig.status_as_on (string, yyyy-MM-dd) = the migration date, e.g. 2026-10-30. "
                 "Fixed date = reruns update instead of adding reports.",
                 "New transform map 'Import Adaptive Status Reports' : u_import_adaptive_projects -> project_status, "
                 "Order 200 (runs after the project map, Order 100).",
                 "Field map 1: [Script] -> project, Coalesce = true: answer = sys_id of pm_project whose correlation_id = "
                 "source.u_correlation_id.",
                 "Field map 2: [Script] -> as_on, Coalesce = true: answer = gs.getProperty('x_mig.status_as_on').",
                 "onBefore (below) sets the health fields and comments; rows with no health at all are skipped.",
                 "Run Transform: select BOTH maps (projects, then status).",
                 "Done when each pilot project has 1 status report whose colours match Adaptive, and a rerun adds none.",
             ]),
             (CODE, "(function runTransformScript(source, map, log, target) {\n"
                    "    var rag = { 'on target': 'green', 'on plan': 'green', 'needs attention': 'yellow', 'in trouble': 'red' };\n"
                    "    function c(v) { return rag[(v + '').toLowerCase().trim()] || ''; }\n"
                    "    target.overall_health = c(source.u_health);\n"
                    "    target.schedule  = c(source.u_schedule_health);\n"
                    "    target.cost      = c(source.u_budget_health);\n"
                    "    target.resources = c(source.u_resource_health);\n"
                    "    target.scope     = c(source.u_scope_health);\n"
                    "    if (!target.overall_health && !target.schedule && !target.cost && !target.resources && !target.scope) { ignore = true; return; }\n"
                    "    var risk = source.u_risk_health + '';\n"
                    "    var notes = (source.u_update_notes + '').replace(/<[^>]+>/g, ' ').replace(/&nbsp;/g, ' ');\n"
                    "    target.comments = (risk ? 'Risk health (Adaptive): ' + risk + '\\n' : '') + notes;\n"
                    "})(source, map, log, target);"),
             (TABLE, ["Adaptive value", "project_status value"], [
                 ["On Target / On Plan", "green"], ["Needs Attention", "yellow"], ["In Trouble", "red"], ["blank", "(empty)"],
             ]),
         ]),
    dict(num="STRY0066702", pts=3, wave="C", title="Resource assignments: who works on each task",
         body=[
             (P, SHARED),
             (P, "Scope: put the people on the tasks (assigned_to + additional_assignee_list). Full Resource "
                 "Management (resource plans / allocations) is NOT in this story (decision D2). Rate models are "
                 "removed unless time is loaded as time cards (decision D3)."),
             (UL, [
                 "Adaptive: Data Source 'Adaptive Resources - API' (STRY0067506) loads one row per resource link with "
                 "Project Correlation ID = the work item (task or project) and Resource Email.",
                 "Transform map u_import_adaptive_resources -> pm_project_task, coalesce on correlation_id = the "
                 "task's correlation id (field map [Script] -> correlation_id returning source.u_project_correlation_id).",
                 "onBefore: first resource on a task -> assigned_to; further resources are added to "
                 "additional_assignee_list (no duplicates). Users matched on email; no match -> logged, left blank.",
                 "Links whose work item is a project (ADAPTIVE:/Project/...) update additional_assignee_list on pm_project.",
                 "Asana: no separate feed - assignee and followers already come with the Work Plan (NEW-11/NEW-12).",
                 "Done when the pilot tasks show the same people as Adaptive.",
             ]),
         ]),
    dict(num="NEW-13", pts=3, wave="C", title="Time entries: actual effort per task (Option A, recommended)",
         body=[
             (P, SHARED),
             (P, "Decision D3. Option A (this story): carry total hours already worked onto each task, so actuals "
                 "are not lost. Option B (time cards per user per week, with rate models) adds cost records, "
                 "approvals and fiscal periods; only if the client needs timesheet history in ServiceNow."),
             (UL, [
                 "Adaptive: Data Source 'Adaptive Time - API' (STRY0067506) on entity [CONFIRM] (e.g. Timesheet), "
                 "fields [CONFIRM] (reported by, work item, date, duration).",
                 "Its script sums hours per work item before inserting: one staging row per task with Project Correlation "
                 "ID (= the task) and Total Hours.",
                 "Transform -> pm_project_task, coalesce via [Script] -> correlation_id = source.u_project_correlation_id; "
                 "Total Hours -> the task's Actual effort field (dev confirms the field name in the pm_project_task dictionary).",
                 "Asana: only if Discovery shows actual_time_minutes; then the Work Plan script also fills Total Hours "
                 "(minutes / 60).",
                 "Done when pilot task actual effort equals the Adaptive totals (spot-check 5 tasks).",
             ]),
         ]),
    dict(num="STRY0067099", pts=3, wave="C", title="Lessons learned: Adaptive -> u_lessons_learned",
         body=[
             (P, SHARED),
             (P, "Before starting: the implementation team confirms u_lessons_learned has correlation_id."),
             (UL, [
                 "Data Source 'Adaptive Lessons - API' (STRY0067506 preset): ENTITY = C_PMOLessonsLearned, "
                 "PROJECT_FIELD = C_RelatedWorkItem.",
                 "Transform map -> u_lessons_learned, coalesce u_correlation_id -> correlation_id.",
                 "u_project: [Script] = sys_id of pm_project whose correlation_id = source.u_project_correlation_id.",
                 "Name -> u_subject; Lesson Type -> u_lesson_type; Topic -> u_topic; Lesson -> u_lesson; ID -> u_name "
                 "(only if u_name is not auto-numbered). Choice values arrive already flattened ('Business Process').",
                 "Done when the pilot projects' lessons show on the project, and a rerun gives 0 inserts.",
             ]),
         ]),
    dict(num="STRY0066721", pts=5, wave="D (blocked)", title="Budgets -> cost plans (cost_plan)",
         body=[
             (P, SHARED),
             (P, "Blocked by decision D4: where budgets live in Adaptive (project budget fields or a budget/expense "
                 "entity - see Discovery), and whether cost plans or demand totals are wanted (not both: cost plans roll "
                 "up and overwrite totals)."),
             (UL, [
                 "Data Source from the confirmed Adaptive source, one row per project per expense type (Capex/Opex), "
                 "Correlation ID = 'ADAPTIVE:' + project id + ':' + expense type.",
                 "Transform -> cost_plan, coalesce on correlation id; task = pm_project via Project Correlation ID.",
                 "Expense type Capex/Opex and resource type map in field map scripts (dev confirms field names in the "
                 "cost_plan dictionary).",
                 "start_date / end_date = project start / end, so ServiceNow builds cost_plan_breakdown per fiscal period. "
                 "Fiscal periods must exist for those dates.",
                 "Done when pilot project totals in ServiceNow equal Adaptive.",
             ]),
         ]),
    dict(num="STRY0066723", pts=5, wave="E - last", title="Attachments: copy files from Asana and Adaptive via the API",
         body=[
             (P, SHARED),
             (P, "Replaces the manual download/upload. Runs LAST, after every record exists. "
                 "Script: servicenow/attachments_fix_script.js (Fix Script)."),
             (UL, [
                 "Asana: GET /attachments?parent=<gid> per migrated record. Files hosted by Asana are downloaded from "
                 "download_url (short-lived signed link, fetched without the token) and saved to the record "
                 "(RESTMessageV2.saveResponseBodyAsAttachment). Google Drive / Box / Dropbox / OneDrive files cannot be "
                 "downloaded: their link is added as a work note.",
                 "Adaptive: files linked to the work item, query and download path [CONFIRM] from Discovery "
                 "(set ADAPTIVE_LINK_QUERY and ADAPTIVE_DOWNLOAD at the top of the script).",
                 "Tables: pm_project, pm_project_task, dmn_demand (+ risk, issue once loaded).",
                 "Run 1 with DRY_RUN = true: System Logs (source x_mig.attachments) lists what would be copied. "
                 "Then DRY_RUN = false.",
                 "Re-run safe: a file with the same name already on the record is skipped. Files over MAX_MB (50) are "
                 "logged for manual copy.",
                 "Done when the final log line shows failed = 0 (or each failure is listed and handled), and a "
                 "spot-check of 10 files opens correctly.",
             ]),
         ]),
]

RUN_ORDER = [
    ("1", "Adaptive Projects - API", "Import Adaptive Projects Transform Map + Import Adaptive Status Reports", ""),
    ("2", "Asana Portfolio - API", "Import Asana HR Project Portfolio Transform Map", ""),
    ("3", "Adaptive Tasks - API, then Adaptive Milestones - API", "Import Adaptive Project Tasks Transform Map", "Transform twice"),
    ("4", "Asana Work Plan - API", "Import Asana Project Tasks Transform Map", "Transform twice"),
    ("5", "Adaptive Resources - API", "Resources -> pm_project_task", ""),
    ("6", "Adaptive Time - API", "Time -> pm_project_task (actual effort)", ""),
    ("7", "Adaptive Risks - API, Adaptive Issues - API", "Existing Risks / Issues maps", ""),
    ("8", "Adaptive Lessons - API", "Lessons -> u_lessons_learned", ""),
    ("9", "(cost plans, when unblocked)", "Budgets -> cost_plan", ""),
    ("10", "attachments_fix_script.js", "(Fix Script, DRY_RUN first)", "Always last"),
]

DECISIONS = [
    ("D1", "Adaptive task states: Draft/Requested/On Hold -> Pending, Active -> Work in Progress, Completed -> Closed "
           "Complete, Cancelled -> Closed Incomplete. Approve or change.", "NEW-12"),
    ("D2", "Full Resource Management (resource plans/allocations) in scope, or task assignees only?", "STRY0066702"),
    ("D3", "Time: Option A (actual effort totals per task) or Option B (weekly time cards + rate models)?", "NEW-13"),
    ("D4", "Budgets: where they live in Adaptive, and cost plans vs demand totals.", "STRY0066721"),
    ("D5", "Status report date (x_mig.status_as_on): migration date?", "STRY0066701"),
    ("D6", "Planned Finish blank: use Adaptive DueDate instead?", "Adaptive Projects - API"),
]


def to_text(body):
    out = []
    for kind, *rest in body:
        if kind in (P, CODE):
            out.append(rest[0])
        elif kind == UL:
            out.extend("- " + i for i in rest[0])
        else:
            out.append(" | ".join(rest[0]))
            out.extend(" | ".join(r) for r in rest[1])
    return "\n".join(out)


def to_html(body):
    e = html.escape
    out = []
    for kind, *rest in body:
        if kind == P:
            out.append(f"<p>{e(rest[0])}</p>")
        elif kind == CODE:
            out.append(f"<pre><code>{e(rest[0])}</code></pre>")
        elif kind == UL:
            out.append('<ul style="list-style-position: inside;">' + "".join(f"<li>{e(i)}</li>" for i in rest[0]) + "</ul>")
        else:
            head = "".join(f"<td><strong>{e(h)}</strong></td>" for h in rest[0])
            rows = "".join("<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>" for r in rest[1])
            out.append(f'<table style="width: 100%;"><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>')
    return "\n".join(out)


def sheet(wb, title, head, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


def main():
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "Stories", ["Number", "Wave", "Short description", "Points", "Acceptance criteria",
                          "Acceptance criteria (HTML - paste into ServiceNow)"],
          [(s["num"], s["wave"], s["title"], s["pts"], to_text(s["body"]), to_html(s["body"])) for s in STORIES],
          [13, 12, 40, 7, 100, 40])
    sheet(wb, "Run order", ["Step", "Data Source", "Transform map", "Note"], RUN_ORDER, [6, 45, 55, 18])
    sheet(wb, "Decisions", ["#", "Decision (BA)", "Story"], DECISIONS, [6, 100, 22])
    wb.save(OUT)
    print(OUT, len(STORIES), "stories,", sum(s["pts"] for s in STORIES), "points")


if __name__ == "__main__":
    main()
