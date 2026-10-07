"""Build docs/rm_story_revised.xlsx: the SPM Data Migration stories (rm_story_2) with review fixes applied.

Sheet "Stories": one row per story, in build order, with plain-text acceptance criteria
plus an HTML copy (last column) to paste into the ServiceNow story.
Sheet "Decisions needed": the open items the stories point to as [BA TO CONFIRM].
"""
import datetime as dt
import html
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "rm_story_revised.xlsx"
ORIGINAL = Path("/root/.claude/uploads/842770d7-6f42-572a-a182-ce07b5e1ad66/12ffe463-rm_story_2.xlsx")

P, UL, TABLE = "p", "ul", "table"

ADAPTIVE_RULES = (
    "Shared Adaptive rules (set in STRY0067506, used by every Adaptive story): "
    "correlation_id = \"ADAPTIVE:\" + the record's API id (e.g. ADAPTIVE:/Project/abc123); "
    "parent_correlation_id is built the same way from the parent's id; "
    "staging columns arrive as u_<name>; all value conversion happens in the transform map."
)

STORIES = [
    # ------------------------------------------------------------------ WAVE 1
    dict(num="STRY0067504", wave="1", pts=4,
         title="Create Outbound REST Message for AdaptiveWork (Clarizen) API",
         changed="Base URL comes from getServerDefinition (not hard-coded api2). Added describeEntities method. "
                 "Allowed a temporary Dev key so work starts before the service account key arrives.",
         body=[
             (P, "This story is complete when:"),
             (UL, [
                 "An Outbound REST Message \"AdaptiveWork API\" is created. Its base URL is the one returned by "
                 "https://api.clarizen.com/v2.0/services/authentication/getServerDefinition (Adaptive tenants sit on "
                 "different data centers, so do not hard-code api2).",
                 "Two POST methods exist: /data/query (CZQL queries) and /metadata/describeEntities (field names).",
                 "Header: Authorization: ApiKey <key>. The key for svc_sn_migration is stored in the Credentials "
                 "module or an encrypted (password2) system property. No keys in scripts.",
                 "If the client approves, a temporary key may be used in Dev until the service account key is issued. "
                 "Swapping keys changes only the credential record.",
                 "Test: a small query (paging limit 5) returns 200 OK and JSON.",
                 "Test: describeEntities for Project, Task and Milestone runs, and the output is attached to this "
                 "story (the dev uses it for the API field names).",
             ]),
         ]),
    dict(num="STRY0067506", wave="1", pts=8,
         title="Create Data Sources and Staging Tables for AdaptiveWork Projects, Demands, Tasks and Milestones",
         changed="Cut to wave 1 entities (rest moved to NEW-03). Added paging, scope filter, one correlation_id rule, "
                 "user emails, milestones. Script only flattens JSON; no value conversion.",
         body=[
             (P, "This story is complete when:"),
             (UL, [
                 "One Data Source (type Custom - Load by Script) exists per staging table: u_imp_adaptive_project, "
                 "u_imp_adaptive_demand, u_imp_adaptive_task.",
                 "Each script calls the AdaptiveWork API REST Message (STRY0067504) with a CZQL query and keeps "
                 "requesting pages while paging.hasMore is true.",
                 "Project query scope: State in [BA TO CONFIRM - in-flight states, see D1].",
                 "Demand vs Project split: [BA TO CONFIRM - see D2]. Until it is given, build Project and Task only.",
                 "Task data source queries Task and Milestone. Milestone rows set milestone = true.",
                 "The script flattens JSON only. It does not convert values: reference objects become the last part "
                 "of the id (\"/State/Active\" -> \"Active\"); people come from the dot-walked email in the CZQL "
                 "SELECT (e.g. ProjectManager.Email).",
                 "Every row gets correlation_id = \"ADAPTIVE:\" + the record's API id (e.g. ADAPTIVE:/Project/abc123) "
                 "and correlation_display = \"Adaptive Work\". Task rows also get parent_correlation_id (parent's id, "
                 "same rule) and project_correlation_id. All Adaptive stories use this rule.",
                 "Running each Data Source loads the staging table without errors, and the staging row count equals "
                 "the Adaptive record count (both written to the log).",
             ]),
         ]),
    dict(num="STRY0066697", wave="1", pts=7,
         title="Create Transform Map: Migrate Active Adaptive Records to Projects (pm_project)",
         changed="Conversion now in the transform map (no upstream pipeline). Business Sponsor -> u_business_owner; "
                 "u_executive_sponsor left empty (no Adaptive source). Start date from Adaptive StartDate, not Project Assigned. Added u_sites script and state value table.",
         body=[
             (P, "Reference the \"Project Mapping\" tab in the attached Adaptive_to_ServiceNow_Mapping_Guide.xlsx for "
                 "field-level directives. \"Project fields for data migration.xlsx\" shows an example Adaptive project."),
             (P, ADAPTIVE_RULES),
             (P, "This story is complete when:"),
             (UL, [
                 "A Transform Map routes u_imp_adaptive_project to pm_project. Coalesce on correlation_id only.",
                 "Core fields map 1:1: Name, Project Manager, Percent Complete. Planned Finish -> end_date.",
                 "start_date = Adaptive StartDate (system field; the Data Source adds it to the CZQL SELECT). If blank, use "
                 "Project Assigned and log it. Project Assigned is not mapped as the start date otherwise.",
                 "Business Sponsor (free text, may list several names) -> u_business_owner: if it holds one name, match it "
                 "to sys_user; if it holds several names or there is no match, leave the field empty and add the full "
                 "text to work_notes. u_executive_sponsor is not mapped (Adaptive has no executive sponsor field).",
                 "Project Manager matches sys_user on email. "
                 "Choice action = Ignore (never create users).",
                 "State and Phase: Choice action = Reject. A field map script converts Adaptive values using the value "
                 "table below.",
                 "Custom fields u_cn, u_sites, u_funding_type, u_funding_source, "
                 "u_business_category map exactly as in the Demand map (STRY0066695).",
                 "u_sites: a field map script splits the site names and returns sys_ids from [BA TO CONFIRM - Sites "
                 "table, see D3]. Unknown names are logged.",
                 "Running the same data a second time in sub-production inserts 0 new records.",
             ]),
             (TABLE, ["Adaptive State", "ServiceNow state"], [
                 ["Draft", "[BA TO CONFIRM - see D4]"],
                 ["Requested", "[BA TO CONFIRM - see D4]"],
                 ["Active", "[BA TO CONFIRM - see D4]"],
                 ["On Hold", "[BA TO CONFIRM - see D4]"],
             ]),
         ]),
    dict(num="STRY0066699", wave="1", pts=7,
         title="Create Transform Map: Migrate Adaptive Work Plans to Project Tasks (pm_project_task)",
         changed="Milestones come from the Milestone entity. Added load order (parents first) and error logging. "
                 "Added state value table.",
         body=[
             (P, "Reference the \"Project Task Mapping\" tab in the attached Adaptive_to_ServiceNow_Mapping_Guide.xlsx."),
             (P, ADAPTIVE_RULES),
             (P, "This story is complete when:"),
             (UL, [
                 "A Transform Map routes u_imp_adaptive_task to pm_project_task. Coalesce on correlation_id only.",
                 "Task Name, Start/End Dates, Assigned To (email, Choice action = Ignore), Percent Complete and "
                 "milestone map 1:1.",
                 "An onBefore script sets target.parent = the record whose correlation_id = u_parent_correlation_id "
                 "(project or task) and target.top_task = the project whose correlation_id = u_project_correlation_id. "
                 "If either is not found, the row errors with a log message.",
                 "Parents load before children (rows sorted by WBS level), or the transform is run a second time.",
                 "State: Choice action = Reject; a field map script converts Adaptive task states using the value table "
                 "[BA TO CONFIRM - see D4].",
                 "Run business rules is on, so dates and percent complete roll up to the project.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
         ]),
    dict(num="STRY0066695", wave="1 (after D2)", pts=7,
         title="Create Transform Map: Adaptive Records to Demands (dmn_demand)",
         changed="Conversion now in the transform map. Total Budget no longer maps to financial_return (that is "
                 "expected return). Blocked until the Demand rule (D2) is known.",
         body=[
             (P, "Reference the \"Demand Mapping\" tab in the attached Adaptive_to_ServiceNow_Mapping_Guide.xlsx."),
             (P, ADAPTIVE_RULES),
             (P, "Blocked until the Demand vs Project rule is confirmed (D2)."),
             (P, "This story is complete when:"),
             (UL, [
                 "A Transform Map routes u_imp_adaptive_demand to dmn_demand. Coalesce on correlation_id only.",
                 "Name, Phase/State and Dates map 1:1 (start date and Business Sponsor follow the Project map rules). Choice action = Reject; a field map script converts Adaptive "
                 "values to ServiceNow choices.",
                 "Capex Budget, Opex Budget and Total Budget map to the demand's matching budget fields (dev confirms "
                 "the field names in the dictionary). Skip this if cost plans are loaded for demands [BA TO CONFIRM - "
                 "see D5].",
                 "Custom fields u_business_owner (from Business Sponsor), u_cn, u_funding_type, u_funding_source, "
                 "u_business_category map 1:1.",
                 "u_sites uses the same field map script as the Project map (STRY0066697).",
                 "User references resolve on email (e.g. BusinessOwner.Email). Choice action = Ignore.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
         ]),
    dict(num="STRY0065263", wave="1 (after Asana re-export)", pts=7,
         title="Asana Data Migration - Field Mapping for Project",
         changed="Aligned to the new HR Project form. Needs an Asana re-export with Task ID. Only top-level, "
                 "in-flight rows load. Removed mappings that would fail (portfolio, primary_goal, impacted business "
                 "units, business category); those fields wait on decision D6.",
         body=[
             (P, "Before starting: the Asana owner re-exports HR_Project_Portfolio_2026 as CSV with the Task ID "
                 "column. Do not open and save it in Excel (it corrupts the 16-digit IDs)."),
             (P, "This story is complete when:"),
             (UL, [
                 "A staging table u_imp_asana_hr_portfolio is created from the CSV.",
                 "A Transform Map routes it to pm_project. Coalesce on correlation_id = \"ASANA:\" + Task ID. "
                 "correlation_display = \"Asana\".",
                 "Only rows with no Parent task load (projects). Rows with a Parent task are workstreams and go to the "
                 "Project Task map (STRY0065275).",
                 "Rows with Project Status Completed or Cancelled are skipped.",
                 "Fields map as in the table below.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
             (TABLE, ["Asana field", "ServiceNow field", "Rule"], [
                 ["Name", "short_description", "Copy as-is."],
                 ["Notes", "description", "Copy as-is. Keep line breaks."],
                 ["Start Date", "start_date", "Date."],
                 ["Due Date", "end_date", "Date."],
                 ["PM Assigned", "project_manager", "Match sys_user by name. N/A, TBD, blank -> empty. Log no-match."],
                 ["Project Status", "state", "In Progress -> Work in Progress; Not Started -> Pending; "
                  "Started/Scoping -> Investigating; Delayed -> Work in Progress; On Hold -> Pending; blank -> Pending."],
                 ["Funding", "u_funding_type", "OpEx -> Opex required; CapEx -> Capex required; blank -> empty."],
             ]),
             (P, "Not mapped until decision D6 is answered (see Asana_HR_Project_to_ServiceNow.xlsx, Questions tab): "
                 "Section/Column (u_functional_area), I&T Integrations Needed (u_it_required), Entity/Organization "
                 "(u_sites), Project Category, HRSP / SME / Project Lead, Assignee, Project Type, Size/Complexity, "
                 "Strategy Focus, Strategy Timeline, HR COEs Engaged, HR Cross-Impacts, Portfolio, Status."),
         ]),
    dict(num="STRY0065275", wave="1 (after Asana re-export)", pts=7,
         title="Asana Data Migration - Field Mapping for Project Task",
         changed="Needs an Asana re-export with Task ID and Assignee Email. Added project link (top_task) and state "
                 "table. Fixed table name planned_task_rel_planned_task. Parent-project fields (COE, Entity) moved "
                 "to decision D6.",
         body=[
             (P, "Before starting: the Asana owner re-exports each task list as an unmodified CSV that includes Task ID "
                 "and Assignee Email. For each file, the BA gives the correlation_id of the ServiceNow project the tasks "
                 "belong to [see D7]."),
             (P, "This story is complete when:"),
             (UL, [
                 "A staging table u_imp_asana_project_tasks is created from the CSV.",
                 "A Transform Map routes it to pm_project_task. Coalesce on correlation_id = \"ASANA:\" + Task ID.",
                 "An onBefore script sets top_task = the project whose correlation_id = u_project_correlation_id, and "
                 "parent = the task with the same name as Parent task in the same project (blank -> the project). "
                 "No match or more than one match: the row errors with a log message.",
                 "Workstream rows from the HR portfolio file (rows with a Parent task) load through this map too.",
                 "Fields map as in the table below.",
                 "After the load, a fix script reads Blocked By / Blocking (task names in the same project) and creates "
                 "planned_task_rel_planned_task records. Unmatched names are logged.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
             (TABLE, ["Asana field", "ServiceNow field", "Rule"], [
                 ["Task Name", "short_description", "Copy as-is."],
                 ["Assigned To", "assigned_to", "Match sys_user on Assignee Email; if blank, on name. Log no-match."],
                 ["Start Date", "start_date", "Date."],
                 ["Due Date", "end_date", "Date."],
                 ["Task Status", "state", "Not Started -> Pending; Started/Scoping -> Investigating; In Progress -> "
                  "Work in Progress; Delayed -> Work in Progress; On Hold -> Pending; Completed -> Closed Complete; "
                  "Cancelled -> Closed Incomplete; blank -> Pending."],
                 ["% Complete", "percent_complete", "0.5 -> 50."],
                 ["Collaborators", "additional_assignee_list", "Names -> sys_user sys_ids. Log no-match."],
                 ["Notes", "work_notes", "Copy as-is."],
                 ["Task Name", "milestone", "true if the name contains \"Milestone\"."],
             ]),
             (P, "Not mapped until decision D6: Section, COE, Tags, Entity/Organization (these describe the project, "
                 "and a transform map cannot set fields on the parent record)."),
         ]),
    dict(num="NEW-01", wave="1", pts=3,
         title="Mock Load 1 in Dev: load and reconcile wave 1",
         changed="New. Gives a clear pass/fail for the first test.",
         body=[
             (P, "This story is complete when:"),
             (UL, [
                 "Wave 1 loads run in Dev in this order: Demand (if D2 is known), Project, Task.",
                 "For each load, a short table records: source count, staging rows, inserted, updated, errors.",
                 "Every error row has a reason in the import log.",
                 "Running every load a second time gives 0 inserts.",
                 "10 projects are spot-checked against Adaptive / Asana: name, manager, dates, state, task count.",
                 "Results are attached to this story.",
             ]),
         ]),
    # ------------------------------------------------------------------ WAVE 2
    dict(num="NEW-03", wave="2", pts=5,
         title="Create Data Sources for remaining AdaptiveWork entities",
         changed="New. Holds the entities removed from STRY0067506 so wave 1 can be tested sooner.",
         body=[
             (P, "Uses the same script pattern and rules as STRY0067506."),
             (P, "This story is complete when:"),
             (UL, [
                 "Data Sources and staging tables exist for: status reports, risks, issues, lessons learned "
                 "(wave 2); budgets, rate models, resource assignments, time (wave 3, only once in scope - see D8).",
                 "Each row carries correlation_id (\"ADAPTIVE:\" + its API id) and parent_correlation_id (the related "
                 "project's or task's id, same rule).",
                 "Staging row count equals the Adaptive record count for each query.",
             ]),
         ]),
    dict(num="STRY0066701", wave="2", pts=5,
         title="Create Transform Map: Migrate Adaptive Health to Status Reports (project_status)",
         changed="Project lookup uses u_parent_correlation_id. RAG values convert in the transform map.",
         body=[
             (P, "Reference the \"Status Report Mapping\" tab in the attached Adaptive_to_ServiceNow_Mapping_Guide.xlsx."),
             (P, ADAPTIVE_RULES),
             (P, "This story is complete when:"),
             (UL, [
                 "A Transform Map routes the status staging table to project_status.",
                 "Coalesce on project (scripted: the pm_project whose correlation_id = u_parent_correlation_id) "
                 "+ as_on (Report Date).",
                 "The 5 RAG fields (Overall, Schedule, Cost, Resource, Scope) map to the OOTB choice fields. Choice "
                 "action = Reject; a field map script converts the Adaptive values to ServiceNow choices "
                 "(value table in the mapping guide).",
                 "Adaptive Update Notes map to comments.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
         ]),
    dict(num="STRY0066718", wave="2", pts=4,
         title="Create Transform Map: Migrate Adaptive Risks to Project Risks (risk)",
         changed="Coalesce on correlation_id (titles repeat). Conversion in the transform map.",
         body=[
             (P, "Reference the \"Project Risk Mapping\" tab in the attached "
                 "Adaptive_to_ServiceNow_RIDAC_Mapping_Guide.xlsx and the example export."),
             (P, ADAPTIVE_RULES),
             (P, "This story is complete when:"),
             (UL, [
                 "A Transform Map routes the risk staging table to risk. Coalesce on correlation_id.",
                 "Parent: the pm_project whose correlation_id = u_parent_correlation_id.",
                 "State converts in a field map script: Draft -> Open; Requested -> Open; Approved -> Work in Progress; "
                 "Resolved -> Closed Complete.",
                 "\"Status\" and \"Update Notes\" map to work_notes so history shows in the activity stream.",
                 "Title, Priority, Impact, Category, Mitigation Plan map 1:1.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
         ]),
    dict(num="STRY0066715", wave="2", pts=4,
         title="Create Transform Map: Migrate Adaptive Issues to Project Issues (issue)",
         changed="Table name consistent (issue). Coalesce on correlation_id (titles repeat). "
                 "Dev confirms Investigating is a valid issue state.",
         body=[
             (P, "Reference the \"Project Issue Mapping\" tab in the attached "
                 "Adaptive_to_ServiceNow_RIDAC_Mapping_Guide.xlsx."),
             (P, ADAPTIVE_RULES),
             (P, "This story is complete when:"),
             (UL, [
                 "A Transform Map routes the issue staging table to issue. Coalesce on correlation_id.",
                 "Parent: the pm_project whose correlation_id = u_parent_correlation_id.",
                 "State converts in a field map script: Draft -> Open; Requested -> Open; Active -> Work in Progress; "
                 "On Hold -> Investigating (dev confirms this choice exists on issue); Completed -> Closed Complete; "
                 "Cancelled -> Closed Incomplete.",
                 "Priority converts: High Priority -> High; Medium Priority -> Moderate; Low Priority -> Low.",
                 "\"Status Update\" and \"Reason for Change\" are joined into description.",
                 "Title, Issue Category, Open Date, Deadline, Date Resolved, Resolution map 1:1.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
         ]),
    dict(num="STRY0067099", wave="2", pts=5,
         title="Create Transform Map: Migrate Adaptive Lessons Learned to Custom Lessons Learned in ServiceNow",
         changed="correlation_id follows the shared rule. Project lookup uses u_parent_correlation_id. "
                 "Fixed the parsing example. Implementation team confirms correlation_id exists on the table.",
         body=[
             (P, ADAPTIVE_RULES),
             (P, "Before starting: the implementation team confirms u_lessons_learned has a correlation_id field."),
             (P, "This story is complete when:"),
             (UL, [
                 "A Transform Map routes the lessons learned staging table to u_lessons_learned. Coalesce on "
                 "correlation_id.",
                 "u_project = the pm_project whose correlation_id = u_parent_correlation_id (built from "
                 "C_RelatedWorkItem.id).",
                 "Fields map as in the table below.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
             (TABLE, ["Adaptive field", "ServiceNow field", "Rule"], [
                 ["id", "correlation_id", "\"ADAPTIVE:\" + id. Coalesce."],
                 ["C_RelatedWorkItem.id", "u_project", "Lookup on pm_project.correlation_id."],
                 ["Name", "u_subject", "Copy as-is."],
                 ["C_LessonType", "u_lesson_type", "Already flattened by the Data Source (last part of the id)."],
                 ["C_Topic", "u_topic", "Already flattened, e.g. \"/C_PMOLessonsLearnedTopic/Business Process\" -> "
                  "\"Business Process\"."],
                 ["C_Lesson", "u_lesson", "Copy as-is (HTML allowed)."],
                 ["SYSID (e.g. C-PMO-7046)", "u_name", "Only if u_name is not auto-numbered."],
             ]),
         ]),
    dict(num="NEW-02", wave="2", pts=3,
         title="Create Transform Map: Migrate Adaptive Task Dependencies (planned_task_rel_planned_task)",
         changed="New. Only if dependencies are in scope (D9).",
         body=[
             (P, ADAPTIVE_RULES),
             (P, "This story is complete when:"),
             (UL, [
                 "The dependency entity is added to the task Data Source pattern (name confirmed via describeEntities).",
                 "A Transform Map routes it to planned_task_rel_planned_task. parent (predecessor) and child (successor) "
                 "are scripted lookups on correlation_id; both are coalesce fields.",
                 "Rows where either side is not found are skipped and logged.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
         ]),
    # ------------------------------------------------------------------ WAVE 3
    dict(num="STRY0066721", wave="3", pts=6,
         title="Create Transform Map: Migrate Adaptive Budgets to Cost Plans (cost_plan)",
         changed="Coalesce on the Adaptive budget line id. Conversion in the transform map. Added fiscal period check. "
                 "Dev confirms field names.",
         body=[
             (P, ADAPTIVE_RULES),
             (P, "Before starting: decision D5 (cost plans vs demand budget totals) and fiscal periods exist in the "
                 "instance for every budget date."),
             (P, "This story is complete when:"),
             (UL, [
                 "A Transform Map routes the budget staging table to cost_plan. Coalesce on correlation_id "
                 "(\"ADAPTIVE:\" + budget line id).",
                 "task = the pm_project or dmn_demand whose correlation_id = u_parent_correlation_id.",
                 "Capex/Opex converts to the cost plan's expense type, and budget categories (Labor, Hardware, "
                 "Software, Vendor) to the resource type, in field map scripts (dev confirms the field names in the "
                 "cost_plan dictionary).",
                 "Planned amount maps to the amount field in the system currency.",
                 "Start and End dates map to start_date and end_date, so OOTB business rules create the "
                 "cost_plan_breakdown records.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
         ]),
    dict(num="STRY0066702", wave="3", pts=5,
         title="Create Transform Map: Migrate Adaptive Rate Models (rate_model_line)",
         changed="Split from the 15-point story into 3 (this one + NEW-04, NEW-05). Waits on scope decision D8.",
         body=[
             (P, "Reference the attached Adaptive_to_ServiceNow_Resource_Mapping_Guide.xlsx."),
             (P, ADAPTIVE_RULES),
             (P, "Before starting: BA confirms historical rates are in scope (D8)."),
             (P, "This story is complete when:"),
             (UL, [
                 "A Transform Map targets rate_model_line, mapping user/role, effective dates and hourly rate. "
                 "Coalesce on correlation_id.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
         ]),
    dict(num="NEW-04", wave="3", pts=5,
         title="Create Transform Map: Migrate Adaptive Resource Assignments",
         changed="Split from STRY0066702.",
         body=[
             (P, ADAPTIVE_RULES),
             (P, "Before starting: BA confirms resource assignments are in scope (D8); dev confirms the target table "
                 "(the original story names resource_assignment)."),
             (P, "This story is complete when:"),
             (UL, [
                 "A Transform Map maps the user (email, Choice action = Ignore) to the project or project task "
                 "(lookup on correlation_id) and Adaptive \"Total Project Assignment\" hours to planned work.",
                 "Coalesce on correlation_id.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
         ]),
    dict(num="NEW-05", wave="3", pts=5,
         title="Create Transform Map: Migrate Adaptive Time Cards (time_card)",
         changed="Split from STRY0066702. Weekly totals are built in the Data Source script (no upstream pipeline).",
         body=[
             (P, ADAPTIVE_RULES),
             (P, "Before starting: BA confirms historical time is in scope (D8). Rate models (STRY0066702) load first."),
             (P, "This story is complete when:"),
             (UL, [
                 "The time Data Source sums daily entries into one row per user, task and week (week starting Sunday).",
                 "A Transform Map targets time_card: Reported By (email), task (lookup on correlation_id), week start, "
                 "total hours.",
                 "Coalesce on user + task + week start.",
                 "Running the same data a second time inserts 0 new records.",
             ]),
         ]),
    dict(num="STRY0066723", wave="Cutover", pts=None,
         title="Manual Extraction and Upload of Legacy Attachments",
         changed="Added a file count first (to point the story) and a list view with correlation_id.",
         body=[
             (P, "This story is complete when:"),
             (UL, [
                 "Files are counted per source first, and the story is pointed from that count.",
                 "Extraction: all relevant attachments are downloaded from Adaptive and Asana and saved in folders "
                 "named by legacy Project/Task ID.",
                 "Record matching: a list view of pm_project, dmn_demand, issue and risk with the correlation_id column "
                 "is used to find each record.",
                 "Upload: files are attached to the matching ServiceNow records.",
                 "Validation: a spot-check confirms files open and names/extensions (.pdf, .docx, .xlsx) are kept.",
             ]),
         ]),
]

DECISIONS = [
    ("D1", "Which Adaptive project States count as in-flight (e.g. Active, On Hold, Draft, Requested)?", "STRY0067506"),
    ("D2", "Which Adaptive records become Demands instead of Projects?", "STRY0067506, STRY0066695"),
    ("D3", "Which table does u_sites reference? (implementation team)", "STRY0066697, STRY0066695"),
    ("D4", "Adaptive project and task State values -> ServiceNow state values.", "STRY0066697, STRY0066699"),
    ("D5", "Demand budgets: map totals on the demand, or load cost plans? Not both (cost plans roll up and "
           "overwrite the totals).", "STRY0066695, STRY0066721"),
    ("D6", "The 12 Asana questions in mapping/Asana_HR_Project_to_ServiceNow.xlsx (Questions tab).",
     "STRY0065263, STRY0065275"),
    ("D7", "Which ServiceNow project does each Asana task file belong to?", "STRY0065275"),
    ("D8", "Are historical rates, resource assignments and time in scope at all?", "STRY0066702, NEW-04, NEW-05"),
    ("D9", "Are Adaptive task dependencies in scope?", "NEW-02"),
]

ASKS = [
    ("Asana owner", "Re-export HR_Project_Portfolio_2026 and each task list as unmodified CSV with Task ID and "
                    "Assignee Email.", "STRY0065263, STRY0065275"),
    ("Client Adaptive admin", "Issue the svc_sn_migration API key.", "STRY0067504"),
    ("Implementation team", "Confirm u_lessons_learned has correlation_id; answer D3.", "STRY0067099"),
]


def to_text(body):
    out = []
    for kind, *rest in body:
        if kind == P:
            out.append(rest[0])
        elif kind == UL:
            out.extend("- " + item for item in rest[0])
        else:
            header, rows = rest
            out.append(" | ".join(header))
            out.extend(" | ".join(r) for r in rows)
    return "\n".join(out)


def to_html(body):
    e = html.escape
    out = []
    for kind, *rest in body:
        if kind == P:
            out.append(f"<p>{e(rest[0])}</p>")
        elif kind == UL:
            out.append('<ul style="list-style-position: inside;">'
                       + "".join(f"<li>{e(i)}</li>" for i in rest[0]) + "</ul>")
        else:
            header, rows = rest
            head = "".join(f"<td><strong>{e(h)}</strong></td>" for h in header)
            body_rows = "".join("<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>" for r in rows)
            out.append(f'<table style="width: 100%;"><thead><tr>{head}</tr></thead><tbody>{body_rows}</tbody></table>')
    return "\n".join(out)


def original_meta():
    """State / Status / Created from the uploaded backlog, keyed by story number."""
    if not ORIGINAL.exists():
        return {}
    ws = load_workbook(ORIGINAL).active
    return {r[0]: dict(state=r[3], status=r[4], created=r[10]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}


def style(ws, widths):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


def main():
    meta = original_meta()
    wb = Workbook()
    ws = wb.active
    ws.title = "Stories"
    ws.append(["Number", "Wave", "Short description", "Acceptance criteria", "What changed", "Points", "State",
               "Status", "Epic", "Created", "Acceptance criteria (HTML - paste into ServiceNow)"])
    for s in STORIES:
        m = meta.get(s["num"], {})
        ws.append([s["num"], s["wave"], s["title"], to_text(s["body"]), s["changed"], s["pts"],
                   m.get("state", "Draft"), m.get("status", "Green"), "SPM Data Migration",
                   m.get("created"), to_html(s["body"])])
        if isinstance(m.get("created"), dt.datetime):
            ws.cell(ws.max_row, 10).number_format = "yyyy-mm-dd"
    style(ws, [13, 12, 40, 90, 45, 7, 14, 8, 18, 11, 40])

    ws = wb.create_sheet("Decisions needed")
    ws.append(["#", "Decision (BA)", "Blocks"])
    for r in DECISIONS:
        ws.append(list(r))
    ws.append([])
    ws.append(["Who", "Ask", "Blocks"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r in ASKS:
        ws.append(list(r))
    style(ws, [20, 100, 35])

    wb.save(OUT)
    print(OUT, len(STORIES), "stories,", sum(s["pts"] or 0 for s in STORIES), "points")


if __name__ == "__main__":
    main()
