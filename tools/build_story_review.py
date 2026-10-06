"""Build docs/Story_Peer_Review.xlsx: peer review of the SPM Data Migration story backlog (rm_story_2)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "docs" / "Story_Peer_Review.xlsx"

# Story | Short name | Pts | Wave | Verdict
PLAN = [
    ("STRY0067504", "Adaptive: Outbound REST Message", 4, "1 - Test now", "Ready after small edits"),
    ("STRY0067506", "Adaptive: Data Sources + staging tables", 8, "1 - Test now", "Rework: cut to Project/Task/Demand first"),
    ("STRY0066697", "Adaptive: Project transform map", 7, "1 - Test now", "Ready after small edits"),
    ("STRY0066699", "Adaptive: Project Task transform map", 7, "1 - Test now", "Ready after small edits"),
    ("STRY0066695", "Adaptive: Demand transform map", 7, "1 - Test now (once Demand rule is known)", "Fix financial fields"),
    ("STRY0065263", "Asana: Project transform map", 7, "1 - Test now", "Rework: conflicts with new HR form"),
    ("STRY0065275", "Asana: Project Task transform map", 7, "1 - Test now", "Blocked: no Task IDs in export"),
    ("NEW-1", "Mock Load 1 in Dev + count check", 3, "1 - Test now", "Add"),
    ("STRY0066718", "Adaptive: Risks", 4, "2", "Small edits"),
    ("STRY0066715", "Adaptive: Issues", 4, "2", "Small edits"),
    ("STRY0067099", "Adaptive: Lessons Learned", 5, "2", "Confirm target fields"),
    ("STRY0066701", "Adaptive: Status reports", 5, "2", "Add to Data Source story"),
    ("NEW-2", "Adaptive: Task dependencies", 3, "2", "Add (if dependencies are in scope)"),
    ("STRY0066721", "Adaptive: Budgets -> cost plans", 6, "3", "Confirm scope + fields"),
    ("STRY0066702", "Adaptive: Rates, resources, time cards", 15, "3", "Split into 3; confirm scope"),
    ("STRY0066723", "Attachments (manual)", None, "3 / cutover", "Point it; confirm volume"),
]

# Severity | Story | Finding | Fix
FINDINGS = [
    # ---- cross-cutting ----
    ("Blocker", "All Adaptive", "Several stories say an 'upstream pipeline' converts values before they reach ServiceNow "
     "(66695, 66697, 66715, 66721, 66702). With the REST pull there is no upstream pipeline: the Data Source "
     "loads raw Adaptive values.",
     "Decide one place for conversions. Recommended: the Data Source script only flattens JSON; every value "
     "conversion (state, phase, priority, RAG, Capex/Opex) lives in the transform map (field map script or onBefore). "
     "Remove 'upstream pipeline' wording from the stories."),
    ("Blocker", "All Adaptive", "The correlation_id format differs between stories: 'Adaptive ID column' (66697), full id "
     "string '/C_PMOLessonsLearned/123' (67099), and the agreed standard 'ADAPTIVE:' + SYSID. API reference fields "
     "(Parent.id, C_RelatedWorkItem.id) return the internal id path, not the SYSID.",
     "One rule, written once in 67506 and referenced everywhere: correlation_id = 'ADAPTIVE:' + API id "
     "(e.g. ADAPTIVE:/Project/abc123). Parent and related-item columns use the same rule, so child lookups "
     "match without a second query."),
    ("Blocker", "67506, 66695", "The Demand vs Project rule is still undefined, so the Demand and Project CZQL queries "
     "cannot be written.", "Get the rule. Until then, test with Projects only."),
    ("Blocker", "67506", "No in-flight filter is stated. Without a WHERE clause the pull returns every project ever "
     "created in Adaptive.", "State the Adaptive State values in scope (e.g. Active, On Hold, Draft?) in the story."),
    ("High", "All", "Several stories point to guides that are not attached here (Adaptive_to_ServiceNow_Mapping_Guide, "
     "RIDAC and Resource mapping guides), so the field lists could not be checked.",
     "Attach them to the stories; send them over if you want them reviewed."),
    # ---- REST ----
    ("Medium", "STRY0067504", "Endpoint hard-codes api2.clarizen.com. Adaptive tenants sit on different data centers.",
     "Use the URL returned by authentication/getServerDefinition. Add a second method, metadata/describeEntities, "
     "so the dev can read real field names."),
    ("Medium", "STRY0067504", "The service account key is not issued yet; this blocks the 200 OK check.",
     "Build and test in Dev with the current key only if the client allows it; then swap the credential record. "
     "Nothing else changes."),
    # ---- Data source ----
    ("High", "STRY0067506", "Covers Projects, Demands, Tasks, Lessons, Risks, Issues, but not Milestones, status reports, "
     "budgets, rates/resources/time or dependencies. Stories 66699 (milestones), 66701, 66721, 66702 have nothing to load.",
     "Wave 1: Project, Task, Milestone (and Demand once the rule exists). Add the rest in a follow-up story with "
     "the wave 2/3 maps."),
    ("High", "STRY0067506", "Paging is not mentioned. /data/query returns pages; without paging only the first page loads.",
     "AC: script loops while paging.hasMore is true; staging row count equals the Adaptive count."),
    ("Medium", "STRY0067506", "People fields: the transform maps resolve users by email, but the API returns user ids.",
     "AC: CZQL selects the dot-walked email (e.g. ProjectManager.Email, C_BusinessOwner.Email) into a flat column."),
    ("Medium", "STRY0067506", "8 points for 6 entities plus parsing.", "Fine once cut to the wave 1 entities."),
    # ---- Project ----
    ("Medium", "STRY0066697", "Custom field list omits u_executive_sponsor, which the Demand map includes. "
     "The story says the two maps must match.", "Add u_executive_sponsor, or remove the 'identical' wording."),
    ("Medium", "STRY0066697", "u_sites needs a script (names -> sys_ids) and the target table for Sites is not stated.",
     "Ask the implementation team which table u_sites references; put it in both 66695 and 66697."),
    ("Low", "STRY0066697", "State values are not listed.", "Add the Adaptive -> ServiceNow state table to the story."),
    # ---- Task ----
    ("Medium", "STRY0066699", "Milestones are a separate Adaptive entity (Milestone), not a flag on Task.",
     "Query Milestone in 67506 and load it through the same map with milestone = true."),
    ("Medium", "STRY0066699", "Parent lookup fails if a child loads before its parent.",
     "AC: rows load parent-first (sort by WBS level), or run the load twice."),
    ("Low", "STRY0066699", "State mapping not listed.", "Add the Adaptive task state table to the story."),
    # ---- Demand ----
    ("High", "STRY0066695", "Total Budget is mapped to financial_return. That field holds the expected return, not a budget.",
     "Confirm the demand budget fields in the dictionary and correct the target."),
    ("Medium", "STRY0066695 / 66721", "If cost plans are also loaded on the demand, ServiceNow rolls them up and overwrites "
     "the directly mapped Capex/Opex totals.", "Pick one: map totals on the demand, or load cost plans. Not both."),
    # ---- Asana project ----
    ("Blocker", "STRY0065263 / 65275", "The Asana CSVs have no Task ID, so correlation_id (the coalesce field) has no source. "
     "Parent task and Blocked By hold task names, which are not unique.",
     "Ask the Asana owner for a re-export with Task ID (standard in Asana CSV exports) before building these maps."),
    ("High", "STRY0065263", "Parent task -> parent project. In the HR portfolio file, rows with a Parent task are the "
     "17 workstreams; they are project tasks, not projects.", "Load top-level rows only as projects; send child rows to the task map."),
    ("High", "STRY0065263", "Section/Column -> primary_portfolio. The new HR form has Functional Area (u_functional_area), "
     "whose choices match 11 of the 18 sections.", "Decide: Section -> u_functional_area (Q1/Q2 in the mapping workbook)."),
    ("High", "STRY0065263", "Entity/Organization -> u_business_category. The new HR form has no u_business_category; "
     "values (BSMH, RSFH, GBS) look like Sites.", "Decide per Q4 in the mapping workbook."),
    ("High", "STRY0065263", "Strategy Timeline -> primary_goal and Strategy Focus -> primary_goal.strategy. Values like "
     "'Long-Term' are not goals, and a transform map cannot set a dot-walked field.", "Decide per Q9; likely description or not migrated."),
    ("High", "STRY0065263", "HR Cross-Impacts -> impacted_business_units. Values (AVP, Skills, Labor...) are not Business Units; "
     "the map would reject or blank them.", "Decide per Q10; likely description."),
    ("Medium", "STRY0065263", "Size/Complexity -> priority with no value table.", "Add the value table, or drop it (Q8)."),
    ("Medium", "STRY0065263", "I&T Integrations goes to description, but the new HR form has u_it_required.", "Map per Q3."),
    ("Medium", "STRY0065263", "New HR form fields missing: u_functional_area, u_project_category, u_it_required.", "Add per the mapping workbook."),
    ("Medium", "STRY0065263", "No scope rule: Completed (13) and Cancelled (2) would load.", "AC: skip Completed and Cancelled."),
    ("Medium", "STRY0065263 / 65275", "Users are resolved 'via email', but the export holds names only.",
     "Resolve by name (sys_user.name), log no-match rows."),
    ("Low", "STRY0065263", "Notes -> work_notes; the mapping uses Detailed Description.", "Pick one."),
    # ---- Asana task ----
    ("High", "STRY0065275", "No column says which project a task belongs to, so top_task cannot be set.",
     "Add: Projects column (or parent chain) -> project correlation_id; onBefore sets parent and top_task "
     "like 66699."),
    ("Low", "STRY0065275", "Table name 'planned_task_rel'.", "planned_task_rel_planned_task."),
    ("Low", "STRY0065275", "Milestone by 'name contains Milestone' is a guess.", "Confirm with the Asana owner, or drop."),
    ("Low", "STRY0065275", "Task Status values not listed.", "Add the agreed task state table."),
    # ---- wave 2 ----
    ("Medium", "STRY0066715 / 66718", "Coalesce on parent + Title. Titles repeat; reruns can duplicate or merge records.",
     "Coalesce on correlation_id like every other map (confirm issue and risk carry it)."),
    ("Medium", "STRY0066715", "Title says project_issue; AC says issue. On Hold -> Investigating may not be an issue state.",
     "Use one table name; check the issue state choices."),
    ("Medium", "STRY0067099", "Coalesce on correlation_id on u_lessons_learned, a custom table.",
     "Confirm with the implementation team that the table has correlation_id."),
    ("Low", "STRY0067099", "Parsing example: input ends in 'Other' but the comment says it returns 'Business Process'.",
     "Fix the example."),
    ("Low", "STRY0066701", "RAG values (Adaptive colours) to ServiceNow choices not listed.", "Add the value table."),
    # ---- wave 3 ----
    ("Medium", "STRY0066721", "Cost plan breakdowns need the fiscal calendar to exist for the budget dates. "
     "'Cost Type + Category' is not unique across periods.", "Confirm fiscal periods exist; coalesce on the Adaptive budget line id."),
    ("Medium", "STRY0066721", "Field names cost_type and resource_type need checking against the instance.",
     "Confirm in the cost_plan dictionary."),
    ("High", "STRY0066702", "15 points, three separate builds. Historical time cards and rates create actual-cost "
     "records and depend on rate models and fiscal periods. resource_assignment needs confirming as the target table.",
     "Confirm historical time/rates are in scope at all. If yes, split into 3 stories."),
    ("Low", "STRY0066723", "No points; finding records by correlation_id needs a list view with that column.",
     "Point it; add correlation_id to the list layout. Count files first to size the effort."),
]

MISSING = [
    ("NEW-1", "Mock Load 1 in Dev",
     "Run wave 1 loads in order (Demand, Project, Task). For each entity record: source count, staging rows, "
     "inserted, updated, errors. Rerun once: 0 inserts. Spot-check 10 projects against Adaptive/Asana."),
    ("NEW-2", "Adaptive task dependencies",
     "Data Source pulls the dependency entity; map to planned_task_rel_planned_task, coalesce on parent + child "
     "(both found by correlation_id). Only if dependencies are in scope."),
    ("Decision", "Value tables",
     "One sheet with every Adaptive and Asana value -> ServiceNow value (state, phase, priority, RAG, funding). "
     "Every map story links to it instead of repeating lists."),
]

FILL = {"Blocker": "F4B6B6", "High": "FCE4D6", "Medium": "FFF2CC", "Low": "EDEDED"}


def write(ws, header, rows, widths, sev_col=None):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for r in rows:
        ws.append(["" if v is None else v for v in r])
        if sev_col is not None and r[sev_col] in FILL:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor=FILL[r[sev_col]])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def main():
    wb = Workbook()
    write(wb.active, ["Story", "What", "Points", "Wave", "Verdict"], PLAN, [14, 40, 8, 30, 40])
    wb.active.title = "Plan"
    write(wb.create_sheet("Findings"), ["Severity", "Story", "Finding", "Fix"], FINDINGS, [10, 20, 70, 70], sev_col=0)
    write(wb.create_sheet("Missing"), ["Id", "Story", "What it covers"], MISSING, [10, 30, 100])
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
