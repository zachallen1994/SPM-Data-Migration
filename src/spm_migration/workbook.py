"""Build mapping/mapping_workbook.xlsx from the mapping CSVs for business workshops."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .common import read_csv

SHEETS = [
    ("Asana to ServiceNow", "asana_to_servicenow.csv"),
    ("Adaptive to ServiceNow", "adaptive_to_servicenow.csv"),
    ("Value Maps", "value_maps.csv"),
]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
STATUS_VALUES = '"Proposed,Confirmed,Deferred,Not migrating,Open question"'


def build(mapping_dir: str | Path, out_path: str | Path) -> Path:
    mapping_dir = Path(mapping_dir)
    wb = Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet("How to use")
    for i, line in enumerate([
        "SPM migration field mapping workbook (generated from mapping/*.csv; edit the CSVs, then rebuild)",
        "",
        "1. Each row maps one source field to one ServiceNow field.",
        "2. In workshops, set 'status' to Confirmed / Deferred / Not migrating / Open question.",
        "3. 'target_table' shows where the value lands; both pm_project and dmn_demand rows exist",
        "   because the same source record can become either, depending on classification.",
        "4. Value Maps holds the state/health/priority translations. target_value must be a stored",
        "   value from your instance's sys_choice export, not the display label.",
        "5. Copy agreed changes back into the CSVs (the CSVs are the source of truth for the code).",
    ], 1):
        readme.cell(row=i, column=1, value=line)
    readme["A1"].font = Font(bold=True, size=13)
    readme.column_dimensions["A"].width = 110

    for title, filename in SHEETS:
        rows = read_csv(mapping_dir / filename)
        if not rows:
            continue
        ws = wb.create_sheet(title)
        headers = list(rows[0].keys())
        ws.append(headers)
        for r in rows:
            ws.append([r.get(h, "") for h in headers])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for idx, h in enumerate(headers, 1):
            width = min(60, max(12, max(len(str(r.get(h, ""))) for r in rows) + 2, len(h) + 2))
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        if "status" in headers:
            col = get_column_letter(headers.index("status") + 1)
            dv = DataValidation(type="list", formula1=STATUS_VALUES, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{len(rows) + 200}")
    out_path = Path(out_path)
    wb.save(out_path)
    return out_path
