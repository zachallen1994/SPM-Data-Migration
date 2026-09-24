"""Build the pipeline's raw files from EXPORT FILES instead of the APIs.

For teams without API access to Asana / Adaptive: the system owners export CSV/Excel,
you drop the files in data/exports/, and these readers write data/raw/<source>/*.jsonl in the
same shape the API extractors produce. Everything downstream (normalize, classify,
transform, load files) is unchanged.

Asana:    one standard CSV/Excel export per Asana project (Project > ... > Export > CSV).
          Needs the "Task ID" column - it is the key that keeps re-exports idempotent.
Adaptive: Excel/CSV list-view exports (Projects, and optionally Tasks, Milestones,
          Dependencies). Column headers are mapped to API field names in settings.yaml.
"""
from __future__ import annotations

import hashlib
import logging
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from .common import iso_date, split_multi, write_csv, write_jsonl

log = logging.getLogger(__name__)

# Columns of a standard Asana CSV export (everything else is a custom field)
ASANA_STANDARD = {
    "task id", "created at", "completed at", "last modified", "name", "section/column", "assignee",
    "assignee email", "start date", "due date", "tags", "notes", "projects", "parent task",
    "parent task id", "blocked by (dependencies)", "blocking (dependencies)", "collaborators",
    "task name", "assigned to", "section",
}
NAME_COLS = ("Name", "Task Name")
ASSIGNEE_COLS = ("Assignee Email", "Assignee", "Assigned To")


def read_table(path: Path) -> list[dict]:
    """Rows of a .csv or .xlsx/.xlsm file (first sheet unless 'file.xlsx#Sheet')."""
    sheet = None
    if "#" in path.name:
        name, sheet = path.name.split("#", 1)
        path = path.with_name(name)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [{h: _cell(v) for h, v in zip(headers, r) if h} for r in rows[1:] if any(v not in (None, "") for v in r)]
    import csv
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return [{(k or "").strip(): (v or "").strip() for k, v in r.items()} for r in csv.DictReader(fh)
                if any((v or "").strip() for v in r.values())]


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _get(row: dict, *names: str) -> str:
    lower = {k.lower(): v for k, v in row.items()}
    for n in names:
        if lower.get(n.lower()):
            return lower[n.lower()]
    return ""


def _synthetic_id(*parts: str) -> str:
    return "x" + hashlib.sha1("|".join(parts).encode()).hexdigest()[:15]


# --------------------------------------------------------------------------- Asana

def import_asana(cfg: dict, export_dir: Path, raw_dir: Path) -> dict[str, int]:
    """cfg = settings.asana.exports: list of {file, role: tracker|plan|standard|intake,
    project_name?, tracker_item? (plan only: tracker task name or Task ID)}."""
    projects, tasks, links, warnings = [], [], [], []
    tracker_ids_by_name: dict[str, str] = {}
    specs = cfg.get("exports") or []
    # Trackers first so plan files can resolve their tracker item by name
    for spec in sorted(specs, key=lambda s: 0 if s.get("role") == "tracker" else 1):
        path = export_dir / spec["file"]
        rows = read_table(path)
        project_gid = spec.get("project_gid") or f"file-{re.sub(r'[^A-Za-z0-9]+', '-', Path(spec['file'].split('#')[0]).stem)}"
        role = spec.get("role", "standard")
        projects.append({"gid": project_gid, "name": spec.get("project_name") or Path(spec["file"]).stem,
                         "_role": role, "archived": False, "completed": False, "custom_fields": [],
                         "created_at": "", "modified_at": ""})
        has_ids = bool(rows) and any(_get(r, "Task ID") for r in rows)
        if not has_ids:
            warnings.append(f"{spec['file']}: no 'Task ID' column - IDs are synthetic and will NOT match "
                            "between mock and final exports. Re-export with Task ID.")
        id_by_name: dict[str, str] = {}
        built = []
        for i, r in enumerate(rows):
            gid = _get(r, "Task ID") or _synthetic_id(spec["file"], _get(r, *NAME_COLS), str(i))
            name = _get(r, *NAME_COLS)
            if name in id_by_name and not _get(r, "Parent task", "Parent task ID"):
                warnings.append(f"{spec['file']}: duplicate top-level task name '{name}' - parent links by name are ambiguous")
            id_by_name.setdefault(name.strip(), gid)
            built.append((gid, r))
        for gid, r in built:
            parent_ref = _get(r, "Parent task ID") or _get(r, "Parent task")
            parent_gid = parent_ref if parent_ref in dict(built) else id_by_name.get(parent_ref.strip(), "")
            if parent_ref and not parent_gid:
                warnings.append(f"{spec['file']}: parent '{parent_ref}' of '{_get(r, *NAME_COLS)}' not found in file")
            section = _get(r, "Section/Column", "Section")
            custom = [{"name": k, "display_value": v} for k, v in r.items()
                      if k and k.strip().lower() not in ASANA_STANDARD and v not in ("", None)]
            assignee = _get(r, *ASSIGNEE_COLS)
            deps = []
            for ref in split_multi(_get(r, "Blocked By (Dependencies)")):
                dep = ref if ref in dict(built) else id_by_name.get(ref, "")
                if dep:
                    deps.append({"gid": dep})
                else:
                    warnings.append(f"{spec['file']}: dependency '{ref}' of '{_get(r, *NAME_COLS)}' not found")
            task = {
                "gid": gid, "_project_gid": project_gid, "name": _get(r, *NAME_COLS), "notes": _get(r, "Notes"),
                "completed": bool(_get(r, "Completed At")), "completed_at": iso_date(_get(r, "Completed At")),
                "created_at": iso_date(_get(r, "Created At")), "modified_at": iso_date(_get(r, "Last Modified")),
                "start_on": iso_date(_get(r, "Start Date")), "due_on": iso_date(_get(r, "Due Date")),
                "assignee": {"email": assignee} if assignee else None,
                "followers": [{"email": c} for c in split_multi(_get(r, "Collaborators"))],
                "tags": [{"name": t} for t in split_multi(_get(r, "Tags"))],
                "memberships": ([{"project": {"gid": project_gid},
                                  "section": {"gid": _synthetic_id(project_gid, section), "name": section}}]
                                if section and not parent_gid else []),
                "dependencies": deps, "custom_fields": custom,
                "permalink_url": f"https://app.asana.com/0/0/{gid}" if not gid.startswith("x") else "",
            }
            if parent_gid:
                task["parent"] = {"gid": parent_gid}
            tasks.append(task)
            if role == "tracker" and not parent_gid:
                tracker_ids_by_name[task["name"].strip()] = gid
        if role == "plan":
            ref = str(spec.get("tracker_item", "")).strip()
            item = ref if ref in tracker_ids_by_name.values() else tracker_ids_by_name.get(ref, "")
            if item:
                links.append({"tracker_item_gid": item, "plan_project_gid": project_gid})
            else:
                warnings.append(f"{spec['file']}: tracker_item '{ref}' not found in the tracker export")
    for w in warnings:
        log.warning(w)
    write_csv(raw_dir / "plan_links_from_exports.csv", links, ["tracker_item_gid", "plan_project_gid"])
    write_csv(raw_dir / "export_warnings.csv", [{"warning": w} for w in warnings], ["warning"])
    for name in ("portfolio_items.jsonl", "status_updates.jsonl"):
        write_jsonl(raw_dir / name, [])
    return {"projects": write_jsonl(raw_dir / "projects.jsonl", projects),
            "tasks": write_jsonl(raw_dir / "tasks.jsonl", tasks), "plan_links": len(links),
            "warnings": len(warnings)}


# --------------------------------------------------------------------------- Adaptive

def _nest(record: dict, path: str, value: str) -> None:
    """'ProjectManager.Email' -> {'ProjectManager': {'Email': value}}."""
    parts = path.split(".")
    cur = record
    for p in parts[:-1]:
        cur = cur.setdefault(p, {})
    cur[parts[-1]] = value


def import_adaptive(cfg: dict, export_dir: Path, raw_dir: Path) -> dict[str, int]:
    """cfg = settings.adaptive.exports: {EntityName: {file, id_column, columns: {API field: header},
    references: {API field: EntityName | auto}}}. Reference columns (Project, Parent,
    Predecessor, Successor) must hold IDs such as SYSID; 'auto' finds the entity that owns the ID."""
    counts, warnings = {}, []
    parsed: dict[str, list[tuple[dict, dict]]] = {}
    owner_of: dict[str, str] = {}                     # SYSID -> full id ('/Task/T-12')
    for entity, spec in (cfg.get("exports") or {}).items():
        rows = read_table(export_dir / spec["file"])
        parsed[entity] = []
        id_col = spec.get("id_column", "SYSID")
        for i, r in enumerate(rows):
            key = _get(r, id_col)
            if not key:
                warnings.append(f"{spec['file']} row {i + 2}: no {id_col} - row skipped")
                continue
            owner_of.setdefault(key, f"/{entity}/{key}")
            parsed[entity].append(({"id": f"/{entity}/{key}"}, r))
    for entity, items in parsed.items():
        spec = cfg["exports"][entity]
        refs = spec.get("references") or {}
        out = []
        for rec, r in items:
            for api_field, header in spec["columns"].items():
                value = _get(r, header)
                if value == "":
                    continue
                if api_field in refs:
                    target = owner_of.get(value) if refs[api_field] == "auto" else f"/{refs[api_field]}/{value}"
                    if not target:
                        warnings.append(f"{spec['file']}: {api_field} '{value}' of {rec['id']} not found in any export")
                        continue
                    rec[api_field] = {"id": target}
                else:
                    _nest(rec, api_field, value)
            out.append(rec)
        counts[entity] = write_jsonl(raw_dir / f"{entity}.jsonl", out)
    for w in warnings:
        log.warning(w)
    write_csv(raw_dir / "export_warnings.csv", [{"warning": w} for w in warnings], ["warning"])
    counts["warnings"] = len(warnings)
    return counts
