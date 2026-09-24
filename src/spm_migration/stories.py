"""Generate the ServiceNow user stories (rm_story) and the Migration Standards page.

Input:   config/user_stories.yaml (simple format: one story = one source file -> one target table)
Outputs: mapping/servicenow_user_stories.xlsx  Stories sheet (import-ready for rm_story) + Migration Standards
         docs/09_servicenow_user_stories.md    readable version
         docs/migration_standards.json          feed for tools/standards_docx.js (Word attachment for the epic)
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .common import load_yaml

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=13, bold=True)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")


def _source_label(story: dict) -> str:
    return "Adaptive" if "ADAPTIVE" in story.get("match_key", "") and "ASANA" not in story.get("match_key", "") else "Asana"


def description_text(story: dict) -> str:
    """Plain-text Detailed Description for rm_story.description."""
    lines = [story["story"].strip(), ""]
    for label, key in (("Source", "source"), ("Target", "target"), ("Match key", "match_key"),
                       ("Prerequisite", "prerequisite")):
        if story.get(key) and story[key] != "n/a":
            lines.append(f"{label}: {story[key]}")
    if story.get("mapping"):
        lines += ["", "FIELD MAPPING (source | ServiceNow | rule)"]
        lines += [f"- {s} | {t} | {r}" for s, t, r in story["mapping"]]
    if story.get("description_block"):
        lines += ["", "DESCRIPTION BLOCK (standards section 7): " + ", ".join(story["description_block"])]
    if story.get("rules"):
        lines += ["", "RULES"] + [f"- {r}" for r in story["rules"]]
    if story.get("out_of_scope"):
        lines += ["", "OUT OF SCOPE: " + "; ".join(story["out_of_scope"])]
    lines += ["", "Standards: see 'Migration Standards' attached to the epic."]
    return "\n".join(lines)


def acceptance_text(story: dict) -> str:
    return "\n".join(f"{i}. {a}" for i, a in enumerate(story["acceptance"], 1))


def _priority_value(label: str) -> int | str:
    return int(label[0]) if label and label[0].isdigit() else label


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: dict[str, int]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = HEAD, FILL, WRAP
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font, cell.alignment = FONT, WRAP
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 18)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build(stories_path: str | Path, xlsx_out: str | Path, md_out: str | Path,
          standards_json_out: str | Path) -> tuple[Path, Path, Path]:
    cfg = load_yaml(stories_path)
    std = cfg["standards"]
    stories = cfg["stories"]

    wb = Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet("Read Me")
    notes = [
        (f"{cfg['epic']}: ServiceNow user stories ({len(stories)} stories, "
         f"{sum(s['points'] for s in stories)} points)", TITLE),
        ("", FONT),
        ("How to load into ServiceNow Agile (rm_story):", BOLD),
        ("1. Import the 'Stories' sheet into rm_story (System Import Sets > Load Data).", FONT),
        ("2. Map: short_description, description, acceptance_criteria, story_points, priority (1-4).", FONT),
        (f"   Set epic = '{cfg['epic']}', product = '{cfg['product']}' and assignment group = '{cfg['assignment_group']}'.", FONT),
        ("3. Attach docs/Migration_Standards.docx (or the 'Migration Standards' sheet) to the epic.", FONT),
        ("   Every story refers to it for the shared rules, so each story stays short.", FONT),
        ("4. The 'key' column (MIG-xx) is a working reference; keep it in the title or as a tag if you want it.", FONT),
        ("", FONT),
        ("Load order: MIG-03 -> MIG-04 -> MIG-01 -> MIG-02 -> MIG-05 -> MIG-06 -> MIG-07; then MIG-08 (mocks) and MIG-09 (cutover).", FONT),
    ]
    for i, (text, font) in enumerate(notes, 1):
        readme.cell(row=i, column=1, value=text).font = font
    readme.column_dimensions["A"].width = 120

    _sheet(wb, "Stories",
           ["key", "short_description", "description", "acceptance_criteria", "story_points", "priority",
            "depends_on", "epic"],
           [[s["key"], s["short_description"], description_text(s), acceptance_text(s), s["points"],
             _priority_value(s["priority"]), s.get("depends_on", ""), cfg["epic"]] for s in stories],
           {"key": 9, "short_description": 50, "description": 110, "acceptance_criteria": 70,
            "story_points": 8, "priority": 8, "depends_on": 16, "epic": 20})

    ws = wb.create_sheet("Migration Standards")
    r = 1
    ws.cell(row=r, column=1, value=std["title"]).font = TITLE
    r += 1
    ws.cell(row=r, column=1, value=std["intro"]).font = FONT
    for sec in std["sections"]:
        r += 2
        ws.cell(row=r, column=1, value=sec["heading"]).font = BOLD
        for b in sec["bullets"]:
            r += 1
            c = ws.cell(row=r, column=1, value=f"- {b}")
            c.font, c.alignment = FONT, WRAP
    ws.column_dimensions["A"].width = 130

    xlsx_out = Path(xlsx_out)
    wb.save(xlsx_out)

    md = [f"# 09: ServiceNow user stories ({cfg['epic']})", "",
          f"{len(stories)} stories, {sum(s['points'] for s in stories)} points. Each story covers one source file and one target table.",
          "The shared rules are written once in **Migration Standards** (below, and in `docs/Migration_Standards.docx` for the epic).",
          "The import-ready version is `mapping/servicenow_user_stories.xlsx`.", "",
          "| Key | Story | Points | Depends on |", "|---|---|---|---|"]
    md += [f"| {s['key']} | {s['short_description']} | {s['points']} | {s.get('depends_on') or '-'} |" for s in stories]
    md += ["", f"## {std['title']}", "", std["intro"]]
    for sec in std["sections"]:
        md += ["", f"**{sec['heading']}**", ""] + [f"- {b}" for b in sec["bullets"]]
    for s in stories:
        md += ["", "---", "", f"## {s['key']}: {s['short_description']}",
               f"*{s['points']} points · Priority {s['priority']}" + (f" · Depends on {s['depends_on']}*" if s.get("depends_on") else "*"),
               "", s["story"].strip(), ""]
        for label, key in (("Source", "source"), ("Target", "target"), ("Match key", "match_key"),
                           ("Prerequisite", "prerequisite")):
            if s.get(key) and s[key] != "n/a":
                md.append(f"- **{label}:** {s[key]}")
        if s.get("mapping"):
            md += ["", "| Source | ServiceNow | Rule |", "|---|---|---|"]
            md += [f"| {a} | `{b}` | {c} |" for a, b, c in s["mapping"]]
        if s.get("description_block"):
            md += ["", "**Description block:** " + ", ".join(s["description_block"])]
        if s.get("rules"):
            md += ["", "**Rules**", ""] + [f"- {x}" for x in s["rules"]]
        if s.get("out_of_scope"):
            md += ["", "**Out of scope:** " + "; ".join(s["out_of_scope"])]
        md += ["", "**Acceptance criteria**", ""] + [f"{i}. {a}" for i, a in enumerate(s["acceptance"], 1)]
    md_out = Path(md_out)
    md_out.write_text("\n".join(md) + "\n", encoding="utf-8")

    standards_json_out = Path(standards_json_out)
    standards_json_out.write_text(json.dumps(std, indent=2), encoding="utf-8")
    return xlsx_out, md_out, standards_json_out
